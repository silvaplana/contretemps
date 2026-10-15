"""Ce que l'assistant « voit » d'un fichier de l'onglet Docs : du texte
(Word, Excel, texte brut), ou le fichier lui-même pour ce que Claude lit
directement (PDF, photos). Rien n'est écrit sur le disque."""

from __future__ import annotations

import base64
import csv
import io
import re
import zipfile
from pathlib import Path

# Au-delà, le fichier n'est pas envoyé (coût, et limite de la requête).
TAILLE_MAX_PDF = 10 * 1024 * 1024
TAILLE_MAX_IMAGE = 5 * 1024 * 1024
CARACTERES_MAX = 60_000

TYPES_IMAGE = {"image/jpeg", "image/png", "image/webp", "image/gif"}
TYPES_TEXTE = (".txt", ".md", ".csv", ".json", ".html", ".htm", ".xml")


def _tronquer(texte: str) -> str:
    if len(texte) <= CARACTERES_MAX:
        return texte
    return texte[:CARACTERES_MAX] + f"\n\n[… fichier tronqué : {len(texte)} caractères au total]"


def _docx(chemin: Path) -> str:
    with zipfile.ZipFile(chemin) as archive:
        xml = archive.read("word/document.xml").decode("utf-8", "replace")
    xml = re.sub(r"</w:p>", "\n", xml)
    xml = re.sub(r"<w:tab/>", "\t", xml)
    texte = re.sub(r"<[^>]+>", "", xml)
    return re.sub(r"\n{3,}", "\n\n", texte).strip()


def _xlsx(chemin: Path) -> str:
    from openpyxl import load_workbook

    classeur = load_workbook(chemin, read_only=True, data_only=True)
    parties = []
    for feuille in classeur.worksheets:
        sortie = io.StringIO()
        ecrivain = csv.writer(sortie, delimiter=";")
        for ligne in feuille.iter_rows(values_only=True):
            if any(v is not None for v in ligne):
                ecrivain.writerow(["" if v is None else v for v in ligne])
        parties.append(f"Feuille « {feuille.title} » :\n{sortie.getvalue()}")
    classeur.close()
    return "\n\n".join(parties)


def contenu_pour_claude(chemin: Path, nom: str, mime: str) -> list[dict]:
    """Blocs à mettre dans le résultat de l'outil « lire_fichier »."""
    entete = {"type": "text", "text": f"Fichier « {nom} » :"}
    minuscule = nom.lower()
    taille = chemin.stat().st_size
    try:
        if mime == "application/pdf" or minuscule.endswith(".pdf"):
            if taille > TAILLE_MAX_PDF:
                return [{"type": "text", "text": f"« {nom} » est un PDF trop gros pour être lu ({taille // 1024} Ko)."}]
            donnees = base64.standard_b64encode(chemin.read_bytes()).decode("ascii")
            return [entete, {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": donnees}}]
        if mime in TYPES_IMAGE:
            if taille > TAILLE_MAX_IMAGE:
                return [{"type": "text", "text": f"« {nom} » est une photo trop grosse pour être regardée."}]
            donnees = base64.standard_b64encode(chemin.read_bytes()).decode("ascii")
            return [entete, {"type": "image", "source": {"type": "base64", "media_type": mime, "data": donnees}}]
        if minuscule.endswith(".docx"):
            return [entete, {"type": "text", "text": _tronquer(_docx(chemin)) or "(document vide)"}]
        if minuscule.endswith(".xlsx"):
            return [entete, {"type": "text", "text": _tronquer(_xlsx(chemin)) or "(classeur vide)"}]
        if minuscule.endswith(TYPES_TEXTE) or mime.startswith("text/"):
            return [entete, {"type": "text", "text": _tronquer(chemin.read_text("utf-8", errors="replace"))}]
    except (OSError, zipfile.BadZipFile, KeyError, ValueError) as erreur:
        return [{"type": "text", "text": f"« {nom} » n'a pas pu être lu ({erreur.__class__.__name__})."}]
    return [{"type": "text", "text": f"« {nom} » ({mime}) n'est pas un format que l'assistant sait lire."}]
