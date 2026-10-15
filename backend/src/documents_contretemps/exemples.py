"""Dossiers et fichiers d'exemple de l'onglet Docs (seed de démo, voir
app/seed.py) : de quoi voir tout de suite les dossiers, les types de
fichiers (PDF, Word, Excel, photos), les statuts, un partage à un cours et
un dossier restreint. Les fichiers sont fabriqués ici, rien à stocker dans
le dépôt.

Ne fait rien si l'école a déjà autre chose que ses dossiers système.
"""

from __future__ import annotations

import io
import zipfile
from xml.sax.saxutils import escape

from PIL import Image, ImageDraw, ImageFont
from sqlalchemy import select
from sqlalchemy.orm import Session

from comptes import Compte
from cours.models import Cours
from documents import DocNode, Documents
from documents.ports import Member

from .adapter import _groupes, nom_affiche, personne_id

ORANGE = (216, 114, 42)
CREME = (255, 250, 241)


def _police(taille: int):
    for nom in ("DejaVuSans.ttf", "Arial.ttf", "LiberationSans-Regular.ttf"):
        try:
            return ImageFont.truetype(nom, taille)
        except OSError:
            continue
    return ImageFont.load_default()


def pdf(titre: str, lignes: list[str]) -> bytes:
    """Une page A4 (rendue en image par Pillow, puis enregistrée en PDF)."""
    page = Image.new("RGB", (1240, 1754), "white")
    dessin = ImageDraw.Draw(page)
    dessin.rectangle([0, 0, 1240, 160], fill=ORANGE)
    dessin.text((80, 50), titre, fill="white", font=_police(56))
    y = 240
    for ligne in lignes:
        dessin.text((80, y), ligne, fill=(58, 36, 16), font=_police(34))
        y += 60
    sortie = io.BytesIO()
    page.save(sortie, "PDF", resolution=150)
    return sortie.getvalue()


def docx(titre: str, paragraphes: list[str]) -> bytes:
    """Un document Word minimal (ouvert par Word, LibreOffice, Google Docs)."""
    corps = f'<w:p><w:pPr><w:pStyle w:val="Title"/></w:pPr><w:r><w:t>{escape(titre)}</w:t></w:r></w:p>'
    corps += "".join(f"<w:p><w:r><w:t xml:space=\"preserve\">{escape(p)}</w:t></w:r></w:p>" for p in paragraphes)
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{corps}</w:body></w:document>"
    )
    sortie = io.BytesIO()
    with zipfile.ZipFile(sortie, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            "</Types>",
        )
        z.writestr(
            "_rels/.rels",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
            'Target="word/document.xml"/></Relationships>',
        )
        z.writestr("word/document.xml", document)
    return sortie.getvalue()


def xlsx(feuille: str, lignes: list[list]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    classeur = Workbook()
    ws = classeur.active
    ws.title = feuille
    for ligne in lignes:
        ws.append(ligne)
    for cellule in ws[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="D8722A")
    for colonne in ws.columns:
        ws.column_dimensions[colonne[0].column_letter].width = max(len(str(c.value or "")) for c in colonne) + 4
    sortie = io.BytesIO()
    classeur.save(sortie)
    return sortie.getvalue()


def photo(texte: str, fond: tuple[int, int, int]) -> bytes:
    image = Image.new("RGB", (1200, 800), fond)
    dessin = ImageDraw.Draw(image)
    for i in range(0, 1200, 60):
        dessin.ellipse([i, 560 + (i % 120), i + 140, 700 + (i % 120)], fill=tuple(min(255, c + 40) for c in fond))
    dessin.text((60, 60), texte, fill="white", font=_police(64))
    sortie = io.BytesIO()
    image.save(sortie, "PNG")
    return sortie.getvalue()


def _membre(db: Session, compte: Compte) -> Member:
    return Member(
        person_id=personne_id(db, compte), org_id=compte.ecole_id, display_name=nom_affiche(compte),
        is_admin=True, is_superuser=False, has_personal_space=True, groups=_groupes(db, compte),
    )


def peupler(db: Session, docs: Documents, admin: Compte) -> bool:
    """Crée les exemples dans l'école de `admin`. Vrai s'ils ont été créés."""
    m = _membre(db, admin)
    org = m.org_id
    deja = db.scalar(select(DocNode.id).where(DocNode.org_id == org, DocNode.system_kind.is_(None)).limit(1))
    if deja is not None:
        return False
    racine = docs.browse(db, m, "private", None)["items"]
    permanent = next(e["id"] for e in racine if e["system_kind"] == "permanent")
    saison = next(e for e in racine if e["system_kind"] == "period" and e["period_current"])
    nom_saison = saison["name"].removeprefix("Saison ")

    def fichier(espace, parent, nom, contenu, statut=None):
        f = docs.upload(db, m, espace, parent, nom, io.BytesIO(contenu))["id"]
        if statut:
            docs.update(db, m, f, change_status=True, status=statut)
        return f

    def dossier(espace, parent, nom, statut=None):
        d = docs.create_folder(db, m, espace, parent, nom)["id"]
        if statut:
            docs.update(db, m, d, change_status=True, status=statut)
        return d

    # Permanent : ce qui vaut d'une saison à l'autre.
    fichier("private", permanent, "Statuts de l'association.pdf", pdf("Statuts", [
        "Article 1 : Nom", "Il est fondé entre les adhérents aux présents statuts", "une association régie par la loi du 1er juillet 1901.",
        "", "Article 2 : Objet", "Cette association a pour objet l'enseignement de la danse.",
    ]), "official")
    fichier("private", permanent, "Règlement intérieur.docx", docx("Règlement intérieur", [
        "Les cours commencent à l'heure : merci d'arriver cinq minutes avant.",
        "La tenue de danse est obligatoire.",
        "Les parents attendent les enfants dans le hall.",
    ]), "validated")
    fichier("private", permanent, "Logo.png", photo("Contretemps", ORANGE))
    bureau = dossier("private", permanent, "Bureau")
    docs.set_access(db, m, bureau, True, [])
    fichier("private", bureau, "Comptes 2025-2026.xlsx", xlsx("Comptes", [
        ["Date", "Libellé", "Recette", "Dépense"],
        ["2025-09-15", "Cotisations septembre", 4200, None],
        ["2025-10-02", "Location salle", None, 650],
        ["2025-12-12", "Costumes gala", None, 980],
    ]))

    # Dossier de la saison en cours.
    s = saison["id"]
    ag = dossier("private", s, f"AG {nom_saison}")
    fichier("private", ag, "Convocation AG.docx", docx("Convocation à l'assemblée générale", [
        "Chers adhérents,", "vous êtes conviés à l'assemblée générale de l'association.",
        "Ordre du jour : rapport moral, rapport financier, élection du bureau.",
    ]), "draft")
    fichier("private", ag, "Bilan financier.xlsx", xlsx("Bilan", [
        ["Poste", "Montant"], ["Cotisations", 38400], ["Subventions", 2500], ["Salaires", -24800], ["Salle", -7800],
    ]), "draft")
    fichier("private", s, "Planning des cours.xlsx", xlsx("Planning", [
        ["Jour", "Heure", "Cours", "Salle"], ["Lundi", "17:00", "Éveil", "Salle 1"], ["Mercredi", "14:00", "Jazz enfants", "Salle 2"],
    ]), "validated")
    photos = dossier("private", s, "Photos gala")
    for i, couleur in enumerate([(124, 77, 206), (14, 138, 138), (216, 114, 42)], start=1):
        fichier("private", photos, f"Gala {i}.png", photo(f"Gala · photo {i}", couleur))

    cours = db.scalars(select(Cours).where(Cours.ecole_id == org).order_by(Cours.ordre, Cours.id)).first()
    if cours is not None:
        dossier_cours = dossier("private", s, f"Cours {cours.nom}")
        docs.set_access(db, m, dossier_cours, False, [{"type": "group", "id": f"course:{cours.id}", "level": "view"}])
        fichier("private", dossier_cours, "Chorégraphie du gala.pdf", pdf(f"Cours {cours.nom}", [
            "Chorégraphie du gala", "1. Entrée côté jardin", "2. Ronde", "3. Salut",
        ]))

    # Public : lisible par tout le monde.
    fichier("public", None, f"Tarifs {nom_saison}.pdf", pdf(f"Tarifs {nom_saison}", [
        "Éveil : 220 €", "Jazz enfants : 280 €", "Cotisation : 15 €",
    ]), "official")
    fichier("public", None, "Affiche du spectacle.png", photo("Spectacle de fin d'année", (31, 63, 143)))

    # Personnel de l'admin.
    fichier("personal", None, "Mes notes.docx", docx("Mes notes", ["Penser à réserver la salle pour le gala."]))
    return True
