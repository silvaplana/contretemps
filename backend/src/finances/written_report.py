"""Rapport financier de l'AG (spec/modules/financier.md §8.6), dans la
forme du modèle de la trésorière : introduction, A Recettes, B Dépenses,
points à revoir, conclusion.

Le modèle de langage rédige le texte, **jamais un chiffre** : il écrit des
repères `{{clé}}` que l'appli remplace par les montants et pourcentages
qu'elle a calculés (Synthèse de la saison et de la précédente). Un nombre
qui resterait écrit en toutes lettres par le modèle (« 1 234 € », « 12 % »)
est signalé au trésorier, comme dans l'AG (§5.1). Word (python-docx) et
PDF (pymupdf) ; le texte est gardé dans `fin_written_reports`.
"""

from __future__ import annotations

import html
import io
import json
import re
from datetime import datetime, timezone

import anthropic
from docx import Document
from docx.shared import Pt
from sqlalchemy import select
from sqlalchemy.orm import Session

from .ai import DEFAULT_MODEL, MODELS, AnalysisError, cost_usd
from .exports import _CSS, _fr, html_to_pdf
from .models import FinWrittenReport
from .ports import Member, Season
from .slides import amounts_in

INSTRUCTIONS = """Tu es le trésorier bénévole d'une association et tu rédiges le rapport financier \
présenté à l'assemblée générale, en français simple, au nom de l'association (« nous »).

Tu reçois les chiffres de la saison et de la saison précédente (JSON), avec pour chacun une CLÉ. \
Règle absolue : tu n'écris JAMAIS un montant ni un pourcentage en chiffres. À la place tu écris le \
repère de la clé entre doubles accolades, par exemple « les cotisations ont rapporté {{produits.756}}, \
soit {{ecart.756}} ({{pct.756}}) par rapport à l'an dernier ». L'application remplace chaque repère \
par le chiffre exact. Un repère inconnu est effacé : n'utilise que les clés fournies.

Forme attendue (celle du rapport de l'an dernier) :
- une introduction d'une phrase (période couverte) ;
- A - Recettes : d'où viennent les recettes, ce qui a bougé par rapport à l'an dernier, poste par poste \
quand c'est parlant ;
- B - Dépenses : les postes principaux, ce qui a monté ou baissé, et pourquoi quand les consignes le disent ;
- Points à revoir : seulement ce que les consignes du trésorier demandent (sinon une liste vide) ;
- une conclusion : le résultat de la saison ({{resultat}}, excédent ou déficit) et la trésorerie.
Phrases courtes, pas de jargon comptable, pas de tableau. Ne commente pas un poste à zéro."""

SCHEMA = {
    "type": "object",
    "properties": {
        "introduction": {"type": "string"},
        "recettes": {"type": "array", "items": {"type": "string"}, "description": "Paragraphes de A - Recettes."},
        "depenses": {"type": "array", "items": {"type": "string"}, "description": "Paragraphes de B - Dépenses."},
        "points": {"type": "array", "items": {"type": "string"}, "description": "Points à revoir, un par ligne."},
        "conclusion": {"type": "string"},
    },
    "required": ["introduction", "recettes", "depenses", "points", "conclusion"],
    "additionalProperties": False,
}

_PLACEHOLDER = re.compile(r"\{\{\s*([\w.\-]+)\s*\}\}")
_PERCENT = re.compile(r"\d+(?:[,.]\d+)?\s?%")


def _now() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def _eur(value: float) -> str:
    text = f"{abs(value):,.2f}".replace(",", " ").replace(".", ",")
    return ("−" if value < 0 else "") + text + " €"


def _pct(value: float | None) -> str:
    return "—" if value is None else f"{value:+.1f} %".replace(".", ",")


def figures(current: dict, previous: dict | None) -> dict[str, dict]:
    """Les chiffres que le modèle peut citer, par clé : {clé: {"label",
    "value" (texte prêt à insérer), "raw"}}. Postes de la Synthèse (montant,
    écart et % avec la saison précédente), totaux, résultat, trésorerie,
    événements."""
    out: dict[str, dict] = {}

    def put(key: str, label: str, text: str, raw) -> None:
        out[key] = {"label": label, "value": text, "raw": raw}

    def lines(synthese: dict, side: str) -> dict[str, dict]:
        return {l["number"]: l for group in synthese[side] for l in group["lines"]}

    for side, prefix in (("produits", "produits"), ("charges", "charges")):
        now, before = lines(current, side), lines(previous, side) if previous else {}
        for number, line in now.items():
            put(f"{prefix}.{number}", line["label"], _eur(line["amount"]), line["amount"])
            if number in before and before[number]["amount"]:
                delta = line["amount"] - before[number]["amount"]
                put(f"ecart.{number}", f"Écart {line['label']}", _eur(delta), delta)
                put(f"pct.{number}", f"Variation {line['label']}", _pct(delta / abs(before[number]["amount"]) * 100), delta)
                put(f"precedent.{number}", f"{line['label']} l'an dernier", _eur(before[number]["amount"]), before[number]["amount"])
    for key, label in (("totalProduits", "Total des recettes"), ("totalCharges", "Total des dépenses"), ("result", "Résultat")):
        put({"result": "resultat"}.get(key, key), label, _eur(current[key]), current[key])
        if previous:
            put(f"precedent.{key}", f"{label} l'an dernier", _eur(previous[key]), previous[key])
    put("tresorerie", "Trésorerie", _eur(current["treasuryTotal"]), current["treasuryTotal"])
    put("fournisseursAPayer", "Fournisseurs à payer", _eur(current["unpaidTotal"]), current["unpaidTotal"])
    for e in current["events"]:
        put(f"{e['event']}.ressources", f"Ressources {e['event']}", _eur(e["resources"]), e["resources"])
        put(f"{e['event']}.cout", f"Coût {e['event']}", _eur(e["cost"]), e["cost"])
        put(f"{e['event']}.benefice", f"Bénéfice {e['event']}", _eur(e["profit"]), e["profit"])
    return out


def fill(text: str, available: dict[str, dict]) -> tuple[str, list[str]]:
    """Remplace les repères ; renvoie le texte et les alertes (repère inconnu,
    nombre écrit par le modèle)."""
    warnings: list[str] = []

    def replace(match: re.Match) -> str:
        key = match.group(1)
        if key in available:
            return available[key]["value"]
        warnings.append(f"Repère inconnu « {key} » effacé.")
        return ""

    filled = _PLACEHOLDER.sub(replace, text)
    known = {abs(round(f["raw"])) for f in available.values() if isinstance(f["raw"], (int, float))}
    for value in amounts_in(_PLACEHOLDER.sub("", text)):
        if not any(abs(round(value) - k) <= 1 for k in known):
            warnings.append(f"{_eur(round(value))[:-5]} € écrit par le modèle, à vérifier.")
    for pct in _PERCENT.findall(_PLACEHOLDER.sub("", text)):
        warnings.append(f"« {pct} » écrit par le modèle, à vérifier.")
    return " ".join(filled.split()), warnings


class WrittenReports:
    def __init__(self, synthese, seasons, note_cost=None, client: anthropic.Anthropic | None = None, org_name=None, store=None) -> None:
        """synthese(db, member, season) : la Synthèse (ag_documents) ;
        seasons(db, org_id) ; note_cost(db, org_id, dollars, détail) ;
        org_name(db, org_id) ; store(db, member, chemin, fichiers)."""
        self.synthese = synthese
        self.seasons = seasons
        self.note_cost = note_cost
        self._client = client
        self.org_name = org_name
        self.store = store

    def _anthropic(self) -> anthropic.Anthropic:
        if self._client is None:
            self._client = anthropic.Anthropic()
        return self._client

    # ----- lecture -----

    def get(self, db: Session, report_id: int) -> dict | None:
        row = db.scalar(
            select(FinWrittenReport).where(FinWrittenReport.report_id == report_id).order_by(FinWrittenReport.id.desc())
        )
        return self._dict(row) if row is not None else None

    @staticmethod
    def _dict(row: FinWrittenReport) -> dict:
        content = json.loads(row.content)
        return {
            "id": row.id,
            "prompt": row.prompt,
            "model": row.model,
            "modelLabel": MODELS.get(row.model, {}).get("label", row.model),
            "costUsd": row.cost_usd,
            "createdAt": row.created_at.isoformat(timespec="seconds"),
            **content,
        }

    # ----- rédaction -----

    def write(self, db: Session, member: Member, report_id: int, season: Season, prompt: str, model: str, analysis: list[str]) -> dict:
        """Fait rédiger le rapport, place les chiffres, garde le texte."""
        if model not in MODELS:
            raise AnalysisError(f"Modèle inconnu : {model}")
        seasons = self.seasons(db, member.org_id)
        index = next((i for i, s in enumerate(seasons) if s.id == season.id), None)
        previous = seasons[index - 1] if index else None
        current = self.synthese(db, member, season)
        before = self.synthese(db, member, previous) if previous is not None else None
        available = figures(current, before)
        data = {
            "saison": season.name,
            "periode": {"debut": current["start"], "fin": current["cutoff"]},
            "saisonPrecedente": previous.name if previous else None,
            "chiffres": {key: {"libelle": f["label"], "valeur": f["value"]} for key, f in available.items()},
            "analyseDuBilan": analysis,
        }
        text = "Chiffres de la saison (JSON ; cite-les par leur clé entre doubles accolades) :\n" + json.dumps(data, ensure_ascii=False, indent=1)
        if prompt.strip():
            text += "\n\nConsignes du trésorier :\n" + prompt.strip()
        params = {
            "model": model,
            "max_tokens": 8000,
            "system": INSTRUCTIONS,
            "messages": [{"role": "user", "content": text}],
            "output_config": {"format": {"type": "json_schema", "schema": SCHEMA}},
        }
        if model != "claude-haiku-5-5":
            params["betas"] = ["server-side-fallback-2026-07-01"]
            params["fallbacks"] = "default"
        try:
            with self._anthropic().beta.messages.stream(**params) as stream:
                response = stream.get_final_message()
        except Exception as exc:  # noqa: BLE001 - toute erreur de l'appel en message clair
            raise AnalysisError(f"La rédaction a échoué : {exc}") from exc
        cost = cost_usd(response.model or model, response.usage)
        if response.stop_reason == "refusal":
            raise AnalysisError("La rédaction a été refusée.", cost)
        block = next((b.text for b in response.content if b.type == "text"), None)
        try:
            draft = json.loads(block or "")
        except json.JSONDecodeError as exc:
            raise AnalysisError("Rapport illisible.", cost) from exc
        content = self.compose(draft, available, season, current)
        row = FinWrittenReport(
            report_id=report_id, prompt=prompt.strip(), model=model, served_by=response.model or model, cost_usd=cost,
            content=json.dumps(content, ensure_ascii=False), created_by=member.person_id,
        )
        db.add(row)
        db.commit()
        if cost and self.note_cost is not None:
            self.note_cost(db, member.org_id, cost, f"Rapport financier {season.name}")
        return self._dict(row)

    @staticmethod
    def compose(draft: dict, available: dict[str, dict], season: Season, synthese: dict) -> dict:
        """Le texte final : repères remplacés, alertes rassemblées."""
        warnings: list[str] = []

        def one(text: str) -> str:
            filled, w = fill(str(text), available)
            warnings.extend(w)
            return filled

        def many(items) -> list[str]:
            return [one(t) for t in items if str(t).strip()]

        return {
            "season": season.name,
            "period": {"start": synthese["start"], "end": synthese["cutoff"]},
            "introduction": one(draft.get("introduction", "")),
            "recettes": many(draft.get("recettes", [])),
            "depenses": many(draft.get("depenses", [])),
            "points": many(draft.get("points", [])),
            "conclusion": one(draft.get("conclusion", "")),
            "warnings": warnings,
        }

    # ----- fichiers -----

    def docx(self, content: dict, org_name: str) -> bytes:
        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        doc.add_heading(f"Rapport financier de l'association {org_name}".strip(), level=1)
        doc.add_paragraph(f"Saison {content['season']}, du {_fr(content['period']['start'])} au {_fr(content['period']['end'])}")
        doc.add_paragraph(content["introduction"])
        doc.add_heading("A - Recettes", level=2)
        for p in content["recettes"]:
            doc.add_paragraph(p)
        doc.add_heading("B - Dépenses", level=2)
        for p in content["depenses"]:
            doc.add_paragraph(p)
        if content["points"]:
            doc.add_heading("Points à revoir", level=2)
            for p in content["points"]:
                doc.add_paragraph(p, style="List Bullet")
        doc.add_heading("Conclusion", level=2)
        doc.add_paragraph(content["conclusion"])
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    def pdf(self, content: dict, org_name: str) -> bytes:
        e = html.escape
        parts = [
            f"<h1>Rapport financier de l'association {e(org_name)}</h1>",
            f"<p class='muted'>Saison {e(content['season'])}, du {_fr(content['period']['start'])} au {_fr(content['period']['end'])}</p>",
            f"<p>{e(content['introduction'])}</p>",
            "<h2>A - Recettes</h2>",
            *(f"<p>{e(p)}</p>" for p in content["recettes"]),
            "<h2>B - Dépenses</h2>",
            *(f"<p>{e(p)}</p>" for p in content["depenses"]),
        ]
        if content["points"]:
            parts.append("<h2>Points à revoir</h2><ul>" + "".join(f"<li>{e(p)}</li>" for p in content["points"]) + "</ul>")
        parts.append(f"<h2>Conclusion</h2><p>{e(content['conclusion'])}</p>")
        return html_to_pdf("".join(parts), _CSS + " body { font-size: 11pt; } p { margin: 6pt 0; }", landscape=False)

    def produce(self, db: Session, member: Member, report_id: int, season: Season, fmt: str) -> tuple[bytes, str, str]:
        content = self.get(db, report_id)
        if content is None:
            raise LookupError("Pas encore de rapport : faites-le rédiger d'abord.")
        org_name = self.org_name(db, member.org_id) if self.org_name else ""
        if fmt == "docx":
            data, media = self.docx(content, org_name), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        elif fmt == "pdf":
            data, media = self.pdf(content, org_name), "application/pdf"
        else:
            raise ValueError("Format inconnu (docx ou pdf)")
        name = f"Rapport financier {season.name}.{fmt}"
        if self.store is not None:
            import tempfile
            from pathlib import Path

            with tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / name
                path.write_bytes(data)
                self.store(db, member, ["AG"], [(name, path)])
        return data, name, media
