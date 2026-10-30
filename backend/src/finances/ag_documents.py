"""Documents de l'AG (spec/modules/financier.md §8.6) : Synthèse
Dépenses-Recettes, Balance, Grand livre, Honoraires par professeur, calculés
au centime depuis les lignes de banque (compte du plan confirmé, voir
accounting.py) et la caisse, pour la saison d'un bilan. Jamais un chiffre
écrit par un modèle : des additions.

Comptabilité de trésorerie : chaque ligne de banque est une écriture à deux
lignes (compte du plan contre 512x la banque, 530 la caisse) ; les soldes de
départ des comptes de trésorerie sont repris dans 110 « Report à nouveau »
pour que la Balance soit équilibrée.

Niveau Complet (partie double, §8.7) : une ligne de banque qui a une pièce
passe par le tiers (achat du jour de la facture 6xx contre 401, honoraires et
notes de frais contre 467x ; paiement 401 contre 512x, même lettre) ; les
soldes de départ viennent de la saisie du trésorier (fin_opening_balances)
ou de la clôture de la saison précédente (fin_closings), sinon de la
trésorerie des relevés. Les écritures ne sont pas stockées : elles sont
recalculées depuis les lignes de banque, de caisse et les pièces, donc
toujours justes ; la clôture fige le résultat et les soldes reportés.

Une ligne sans compte confirmé va dans « À classer » (dépense ou recette
selon le sens) plutôt que d'être inventée.
"""

from __future__ import annotations

import html
import io
from collections import defaultdict
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from .accounting import Accounting
from .exports import _CSS, _eur, _fr, html_to_pdf
from .models import EXPENSE, FULL, INCOME, THIRD_PARTY, FinCashLine, FinClosing, FinOpeningBalance
from .ports import Member, Season

# Numéros des comptes de trésorerie et de report (plan associatif).
BANK_BASE = "512"
CASH = "530"
CARRY_FORWARD = "110"
# Compte de tiers par nature de pièce (préfixe du plan associatif).
THIRD_PARTY_PREFIX = {"achat": "401", "honoraires": "467", "frais": "467"}
THIRD_PARTY_JOURNAL = {"achat": "AC", "honoraires": "HO", "frais": "NF"}
UNCLASSIFIED = {EXPENSE: ("65900", "Dépenses à classer"), INCOME: ("75900", "Recettes à classer")}

CLASS_LABELS = {
    "60": "Achats",
    "61": "Services extérieurs",
    "62": "Autres services extérieurs",
    "63": "Impôts et taxes",
    "64": "Charges de personnel",
    "65": "Autres charges",
    "66": "Charges financières",
    "67": "Charges exceptionnelles",
    "68": "Dotations",
    "70": "Ventes et prestations",
    "74": "Subventions",
    "75": "Produits de gestion courante",
    "76": "Produits financiers",
    "77": "Produits exceptionnels",
    "78": "Reprises",
}

KINDS = {
    "synthese": ("Synthèse Dépenses-Recettes", "xlsx", "simple"),
    "balance": ("Balance", "pdf", "complet"),
    "grand-livre": ("Grand livre", "pdf", "complet"),
    "honoraires": ("Honoraires par professeur", "xlsx", "complet"),
    "fec": ("FEC", "txt", "simple"),
}

# Fichier des écritures comptables (article A47 A-1 du livre des procédures
# fiscales) : 18 colonnes séparées par une tabulation, dates AAAAMMJJ,
# montants à virgule.
FEC_COLUMNS = [
    "JournalCode", "JournalLib", "EcritureNum", "EcritureDate", "CompteNum", "CompteLib", "CompAuxNum", "CompAuxLib",
    "PieceRef", "PieceDate", "EcritureLib", "Debit", "Credit", "EcritureLet", "DateLet", "ValidDate", "Montantdevise", "Idevise",
]
JOURNALS = {"AN": "A-nouveaux", "BQ": "Banque", "CA": "Caisse", "AC": "Achats", "HO": "Honoraires", "NF": "Notes de frais"}


@dataclass
class AgPorts:
    """Ce que l'appli hôte fournit : les pièces à payer (fournisseurs,
    Synthèse), les honoraires payés par prof et par mois (DAS-2), et le
    rangement d'un fichier dans Docs › Financier › AG (ids des fichiers)."""

    unpaid: Callable[[Session, int], list[dict]] | None = None
    fees: Callable[[Session, int, str, str], list[dict]] | None = None
    store: Callable[[Session, Member, list[str], list[tuple[str, Path]]], list[int]] | None = None
    folder: Callable[[Session, Member], int | None] | None = None  # le dossier AG de Docs, pour l'ouvrir
    org_name: Callable[[Session, int], str] | None = None


@dataclass
class Line:
    """Une ligne du grand livre : date, journal, libellé, pièce, débit et
    crédit en centimes."""

    date: str
    journal: str
    label: str
    piece: str
    debit: int = 0
    credit: int = 0
    letter: str = ""
    operation_id: int | None = None


@dataclass
class Account:
    number: str
    label: str
    nature: str  # depense, recette, banque, caisse, report
    event: str | None = None
    lines: list[Line] = field(default_factory=list)

    @property
    def debit(self) -> int:
        return sum(l.debit for l in self.lines)

    @property
    def credit(self) -> int:
        return sum(l.credit for l in self.lines)

    @property
    def balance(self) -> int:
        """Solde signé : positif = débiteur."""
        return self.debit - self.credit


class AgDocuments:
    def __init__(
        self,
        ledger: Callable[[Session, Member], dict],
        accounting: Accounting,
        seasons: Callable[[Session, int], list[Season]],
        ports: AgPorts | None = None,
    ) -> None:
        self.ledger = ledger
        self.accounting = accounting
        self.seasons = seasons
        self.ports = ports or AgPorts()

    # ----- le grand livre en mémoire -----

    def books(self, db: Session, member: Member, season: Season, full: bool | None = None) -> dict[str, Account]:
        """Tous les comptes mouvementés sur la saison, avec leurs lignes, les
        soldes de départ et le report à nouveau. full : partie double (par
        défaut selon le niveau de l'organisation)."""
        ledger = self.ledger(db, member)
        rows = [r for r in reversed(ledger["rows"]) if season.start_date <= r["date"] <= season.end_date]
        # Arrêté au dernier jour connu des comptes (jamais avant le début).
        as_of = max((a["asOf"] for a in ledger["accounts"] if a.get("asOf")), default=season.end_date)
        cutoff = max(season.start_date, min(season.end_date, as_of))
        chart = {a.id: a for a in self.accounting.accounts(db, member.org_id)}
        books: dict[str, Account] = {}

        def account(number: str, label: str, nature: str, event: str | None = None) -> Account:
            return books.setdefault(number, Account(number, label, nature, event))

        bank_numbers = {}
        for i, a in enumerate(sorted(ledger["allAccounts"], key=lambda a: a["id"])):
            bank_numbers[a["id"]] = (f"{BANK_BASE}{i + 1}", a["name"])

        # Soldes de départ de la trésorerie (banque : solde de la veille de la
        # saison, ou premier solde connu moins les opérations qui y mènent ;
        # caisse : lignes antérieures) contre 110.
        def effect(r: dict, aid: int, name: str) -> int:
            """Ce qu'une ligne de l'historique fait au compte aid (centimes)."""
            cents = round(r["amount"] * 100)
            if r.get("transfer"):
                return (-cents if r["accountId"] == aid else 0) + (cents if r["transfer"]["to"] == name else 0)
            return cents if r["accountId"] == aid else 0

        if full is None:
            full = self.accounting.level(db, member.org_id) == FULL
        opening_total = 0
        all_rows = list(reversed(ledger["rows"]))
        entered = self._opening_lines(db, member.org_id, season) if full else None
        if entered is not None:
            for number, label, nature, balance in entered:
                account(number, label, nature).lines.append(
                    Line(season.start_date, "AN", "Solde au début de la saison", "", max(balance, 0), max(-balance, 0))
                )
            bank_numbers_for_opening = {}
        else:
            bank_numbers_for_opening = bank_numbers
        for aid, (number, name) in bank_numbers_for_opening.items():
            known = [r for r in all_rows if r["balances"].get(str(aid)) is not None]
            before = [r for r in known if r["date"] < season.start_date]
            if before:
                opening_cents = round(before[-1]["balances"][str(aid)] * 100)
            else:
                first = next((r for r in known if r["date"] >= season.start_date), None)
                if first is None:
                    continue
                until = all_rows[: all_rows.index(first) + 1]
                opening_cents = round(first["balances"][str(aid)] * 100) - sum(effect(r, aid, name) for r in until)
            acc = account(number, name, "banque")
            acc.lines.append(Line(season.start_date, "AN", "Solde au début de la saison", "", max(opening_cents, 0), max(-opening_cents, 0)))
            opening_total += opening_cents
        cash_before = sum(
            l.amount for l in db.scalars(select(FinCashLine).where(FinCashLine.org_id == member.org_id, FinCashLine.date < date.fromisoformat(season.start_date)))
        ) if entered is None else 0
        if cash_before:
            account(CASH, "Caisse", "caisse").lines.append(Line(season.start_date, "AN", "Solde au début de la saison", "", max(cash_before, 0), max(-cash_before, 0)))
            opening_total += cash_before
        if opening_total:
            account(CARRY_FORWARD, "Report à nouveau", "report").lines.append(
                Line(season.start_date, "AN", "Trésorerie au début de la saison", "", max(-opening_total, 0), max(opening_total, 0))
            )

        # Caisse de la saison, et les retraits au distributeur (la ligne de
        # banque va alors dans 530, pas dans un compte du plan).
        cash_lines = list(
            db.scalars(
                select(FinCashLine).where(
                    FinCashLine.org_id == member.org_id,
                    FinCashLine.date >= date.fromisoformat(season.start_date),
                    FinCashLine.date <= date.fromisoformat(season.end_date),
                ).order_by(FinCashLine.date, FinCashLine.id)
            )
        )
        withdrawals = {l.transfer_operation_id for l in cash_lines if l.transfer_operation_id is not None}

        # Lignes de banque : compte du plan (ou à classer) contre 512x ; en
        # Complet, une ligne avec pièce passe par le tiers, avec une lettre.
        confirmed = self.accounting._lines(db, member.org_id)
        third_parties = [a for a in chart.values() if a.nature == THIRD_PARTY and a.active] if full else []
        letters: dict[str, int] = defaultdict(int)
        for r in rows:
            cents = round(r["amount"] * 100)
            if r.get("transfer"):
                # Virement interne : 512a → 512b (le débit porte la ligne).
                src = bank_numbers[r["accountId"]]
                dst = next((n for aid, n in bank_numbers.items() if n[1] == r["transfer"]["to"]), None)
                account(*src, "banque").lines.append(Line(r["date"], "BQ", r["label"], "", 0, cents))
                if dst:
                    account(*dst, "banque").lines.append(Line(r["date"], "BQ", r["label"], "", cents, 0))
                continue
            bank_number, bank_name = bank_numbers[r["accountId"]]
            journal = f"BQ{bank_number[-1]}"
            piece = r["piece"]["title"] if r.get("piece") else ""
            account(bank_number, bank_name, "banque").lines.append(
                Line(r["date"], journal, r["label"], piece, max(cents, 0), max(-cents, 0), operation_id=r["id"])
            )
            if r["id"] in withdrawals:
                account(CASH, "Caisse", "caisse").lines.append(Line(r["date"], journal, r["label"], "", -cents, 0, operation_id=r["id"]))
                continue
            splits = confirmed.get(r["id"])
            parts = [(chart[l.account_id], l.amount) for l in splits] if splits else [(None, cents)]
            third = _third_party(third_parties, r.get("piece")) if cents < 0 and splits else None
            if third is not None:
                # Achat du jour de la facture : 6xx contre le tiers ; paiement : tiers contre 512x.
                letters[third.number] += 1
                letter = _letter(letters[third.number])
                day = max(r["piece"]["date"], season.start_date) if r["piece"].get("date") else r["date"]
                buy = THIRD_PARTY_JOURNAL[r["piece"]["nature"]]
                tiers = account(third.number, third.label, THIRD_PARTY)
                tiers.lines.append(Line(day, buy, r["piece"]["title"], piece, 0, -cents, letter, r["id"]))
                tiers.lines.append(Line(r["date"], journal, r["label"], piece, -cents, 0, letter, r["id"]))
                for acc_row, amount in parts:
                    acc = account(acc_row.number, acc_row.label, acc_row.nature, acc_row.event)
                    acc.lines.append(Line(day, buy, r["piece"]["title"], piece, -amount, 0, operation_id=r["id"]))
                continue
            for acc_row, amount in parts:
                if acc_row is None:
                    number, label = UNCLASSIFIED[EXPENSE if amount < 0 else INCOME]
                    acc = account(number, label, EXPENSE if amount < 0 else INCOME)
                else:
                    acc = account(acc_row.number, acc_row.label, acc_row.nature, acc_row.event)
                acc.lines.append(Line(r["date"], journal, r["label"], piece, max(-amount, 0), max(amount, 0), operation_id=r["id"]))

        # Caisse : compte du plan contre 530 (un retrait est déjà passé par
        # la banque ci-dessus).
        for l in cash_lines:
            if l.transfer_operation_id is not None:
                continue
            day = l.date.isoformat()
            account(CASH, "Caisse", "caisse").lines.append(Line(day, "CA", l.label, "", max(l.amount, 0), max(-l.amount, 0)))
            acc_row = chart.get(l.account_id)
            if acc_row is None:
                number, label = UNCLASSIFIED[EXPENSE if l.amount < 0 else INCOME]
                acc = account(number, label, EXPENSE if l.amount < 0 else INCOME)
            else:
                acc = account(acc_row.number, acc_row.label, acc_row.nature, acc_row.event)
            acc.lines.append(Line(day, "CA", l.label, "", max(-l.amount, 0), max(l.amount, 0)))
        for acc in books.values():
            acc.lines.sort(key=lambda l: l.date)
        books["_cutoff"] = cutoff  # type: ignore[assignment]
        return books

    # ----- soldes de départ et clôture (Complet, §8.7) -----

    def _previous(self, db: Session, org_id: int, season: Season) -> Season | None:
        seasons = sorted(self.seasons(db, org_id), key=lambda s: s.start_date)
        index = next((i for i, s in enumerate(seasons) if s.id == season.id), None)
        return seasons[index - 1] if index else None

    def closing(self, db: Session, org_id: int, season_id: int) -> FinClosing | None:
        """La clôture en vigueur d'une saison (pas rouverte)."""
        return db.scalar(
            select(FinClosing).where(FinClosing.org_id == org_id, FinClosing.season_id == season_id, FinClosing.reopened_at.is_(None))
        )

    def _opening_lines(self, db: Session, org_id: int, season: Season) -> list[tuple[str, str, str, int]] | None:
        """Les soldes de départ (numéro, libellé, nature, solde signé) :
        saisis par le trésorier, sinon reportés de la clôture précédente,
        sinon None (trésorerie des relevés). Toujours équilibrés par 110."""
        entered = list(
            db.scalars(select(FinOpeningBalance).where(FinOpeningBalance.org_id == org_id, FinOpeningBalance.season_id == season.id))
        )
        if entered:
            lines = [(b.number, b.label, _nature_of(b.number, db, self.accounting, org_id), b.balance) for b in entered]
        else:
            previous = self._previous(db, org_id, season)
            closing = self.closing(db, org_id, previous.id) if previous else None
            if closing is None:
                return None
            lines = [(b["number"], b["label"], b["nature"], b["balance"]) for b in closing.balances]
        lines = [l for l in lines if l[3]]
        gap = sum(l[3] for l in lines)
        if gap:
            lines = [l for l in lines if l[0] != CARRY_FORWARD]
            carried = next((l[3] for l in lines if l[0] == CARRY_FORWARD), 0)
            lines.append((CARRY_FORWARD, "Report à nouveau", "report", carried - gap))
        return lines

    def opening(self, db: Session, member: Member, season: Season) -> dict:
        """Les soldes de départ à l'écran : d'où ils viennent, un solde par
        compte du bilan (en euros), et la vérification des 512x contre les
        relevés."""
        entered = self._opening_lines(db, member.org_id, season)
        source = "saisie" if db.scalar(
            select(FinOpeningBalance.id).where(FinOpeningBalance.org_id == member.org_id, FinOpeningBalance.season_id == season.id)
        ) else ("cloture" if entered is not None else "releves")
        # Trésorerie des relevés au début de la saison : sans soldes saisis,
        # c'est ce que books() prend ; sinon la référence de la vérification.
        statements = self._statement_opening(db, member, season)
        if entered is None:
            lines = [(n, label, nature, balance) for n, (label, nature, balance) in statements.items()]
            total = sum(l[3] for l in lines)
            if total:
                lines.append((CARRY_FORWARD, "Report à nouveau", "report", -total))
        else:
            lines = entered
        by_number = {l[0]: l for l in lines}
        for a in self.accounting.accounts(db, member.org_id, active_only=True):
            if a.nature == THIRD_PARTY or a.number.startswith("47"):
                by_number.setdefault(a.number, (a.number, a.label, a.nature, 0))
        for number, (label, nature, _) in statements.items():
            by_number.setdefault(number, (number, label, nature, 0))
        by_number.setdefault(CARRY_FORWARD, (CARRY_FORWARD, "Report à nouveau", "report", 0))
        accounts = [
            {"number": n, "label": label, "nature": nature, "balance": balance / 100}
            for n, label, nature, balance in sorted(by_number.values(), key=lambda l: l[0])
        ]
        check = [
            {"number": n, "label": label, "statements": balance / 100, "entered": by_number.get(n, (0, 0, 0, 0))[3] / 100}
            for n, (label, _, balance) in statements.items()
        ]
        return {"season": season.name, "source": source, "accounts": accounts, "bankCheck": check, "closed": self.closing(db, member.org_id, season.id) is not None}

    def _statement_opening(self, db: Session, member: Member, season: Season) -> dict[str, tuple[str, str, int]]:
        """Trésorerie des relevés et de la caisse à la veille de la saison :
        {numéro: (libellé, nature, centimes)}, comme books() sans saisie."""
        books = self.books(db, member, season, full=False)
        out = {}
        for a in books.values():
            if isinstance(a, Account) and a.nature in ("banque", "caisse"):
                opening = sum(l.debit - l.credit for l in a.lines if l.journal == "AN")
                out[a.number] = (a.label, a.nature, opening)
        return out

    def set_opening(self, db: Session, member: Member, season: Season, accounts: list[dict]) -> dict:
        """« Reprendre ces soldes » : les soldes de départ saisis (euros,
        positif = débiteur) ; la différence va dans 110."""
        if not member.is_treasurer:
            raise PermissionError("Réservé au trésorier")
        if self.closing(db, member.org_id, season.id) is not None:
            raise ValueError("Saison clôturée : rouvrez-la d'abord.")
        for old in db.scalars(select(FinOpeningBalance).where(FinOpeningBalance.org_id == member.org_id, FinOpeningBalance.season_id == season.id)):
            db.delete(old)
        for a in accounts:
            number = str(a.get("number", "")).strip()
            balance = round(float(a.get("balance", 0) or 0) * 100)
            if not number or not balance:
                continue
            db.add(FinOpeningBalance(org_id=member.org_id, season_id=season.id, number=number[:10], label=str(a.get("label", ""))[:80], balance=balance, created_by=member.person_id))
        db.commit()
        return self.opening(db, member, season)

    def close(self, db: Session, member: Member, season: Season) -> dict:
        """Clôture (§8.7) : le résultat va dans 110, les soldes des comptes de
        bilan sont reportés sur la saison suivante, la saison est verrouillée."""
        if not member.is_treasurer:
            raise PermissionError("Réservé au trésorier")
        if self.accounting.level(db, member.org_id) != FULL:
            raise ValueError("La clôture est au niveau Complet (Comptes › ⋯ › Plan comptable).")
        if self.closing(db, member.org_id, season.id) is not None:
            raise ValueError("Saison déjà clôturée.")
        books = self.books(db, member, season)
        books.pop("_cutoff")
        result = sum(a.balance for a in books.values() if a.nature in (EXPENSE, INCOME))
        balances = {a.number: {"number": a.number, "label": a.label, "nature": a.nature, "balance": a.balance} for a in books.values() if a.nature not in (EXPENSE, INCOME)}
        carry = balances.setdefault(CARRY_FORWARD, {"number": CARRY_FORWARD, "label": "Report à nouveau", "nature": "report", "balance": 0})
        carry["balance"] += result
        db.add(
            FinClosing(
                org_id=member.org_id, season_id=season.id, start_date=date.fromisoformat(season.start_date), end_date=date.fromisoformat(season.end_date),
                result=result, balances=[b for b in sorted(balances.values(), key=lambda b: b["number"]) if b["balance"]], closed_by=member.person_id,
            )
        )
        db.commit()
        return self.closing_status(db, member.org_id, season)

    def reopen(self, db: Session, member: Member, season: Season) -> dict:
        if not member.is_treasurer:
            raise PermissionError("Réservé au trésorier")
        closing = self.closing(db, member.org_id, season.id)
        if closing is None:
            raise ValueError("Saison pas clôturée.")
        closing.reopened_by, closing.reopened_at = member.person_id, datetime.now(timezone.utc).replace(tzinfo=None)
        db.commit()
        return self.closing_status(db, member.org_id, season)

    def closing_status(self, db: Session, org_id: int, season: Season) -> dict | None:
        closing = self.closing(db, org_id, season.id)
        if closing is None:
            return None
        return {"closedAt": closing.closed_at.isoformat(timespec="seconds"), "result": closing.result / 100}

    def operation_entries(self, db: Session, member: Member, operation_id: int, day: str) -> list[dict]:
        """Les écritures générées pour une opération (fiche, §8.4) : une
        ligne par compte mouvementé."""
        season = next((s for s in self.seasons(db, member.org_id) if s.start_date <= day <= s.end_date), None)
        if season is None:
            return []
        books = self.books(db, member, season)
        books.pop("_cutoff")
        entries = []
        for a in sorted(books.values(), key=lambda a: a.number):
            for l in a.lines:
                if l.operation_id == operation_id:
                    entries.append({"date": l.date, "journal": l.journal, "account": f"{a.number} {a.label}", "label": l.label, "debit": l.debit / 100, "credit": l.credit / 100, "letter": l.letter})
        entries.sort(key=lambda e: (e["date"], e["journal"]))
        return entries

    # ----- Synthèse -----

    def synthese(self, db: Session, member: Member, season: Season) -> dict:
        """La Synthèse à l'écran (forme du modèle Excel) : charges et produits
        par classe, totaux, résultat, trésorerie, fournisseurs à payer, un
        bloc par événement."""
        books = self.books(db, member, season)
        cutoff = books.pop("_cutoff")
        accounts = [a for a in books.values() if a.nature in (EXPENSE, INCOME)]

        def grouped(nature: str) -> list[dict]:
            classes: dict[str, list] = defaultdict(list)
            for a in sorted(accounts, key=lambda a: a.number):
                if a.nature == nature:
                    amount = a.balance if nature == EXPENSE else -a.balance
                    if amount:
                        classes[a.number[:2]].append({"number": a.number, "label": a.label, "amount": amount / 100})
            return [
                {"code": code, "label": CLASS_LABELS.get(code, "Autres"), "lines": lines, "total": round(sum(l["amount"] for l in lines), 2)}
                for code, lines in sorted(classes.items())
            ]

        charges, produits = grouped(EXPENSE), grouped(INCOME)
        total_charges = round(sum(c["total"] for c in charges), 2)
        total_produits = round(sum(p["total"] for p in produits), 2)
        treasury = [
            {"label": a.label, "amount": a.balance / 100} for a in sorted(books.values(), key=lambda a: a.number) if a.nature in ("banque", "caisse")
        ]
        unpaid = self.ports.unpaid(db, member.org_id) if self.ports.unpaid else []
        events: dict[str, dict] = {}
        for a in accounts:
            if a.event:
                e = events.setdefault(a.event, {"event": a.event, "resources": 0.0, "cost": 0.0})
                if a.nature == INCOME:
                    e["resources"] += -a.balance / 100
                else:
                    e["cost"] += a.balance / 100
        for e in events.values():
            e["profit"] = round(e["resources"] - e["cost"], 2)
        return {
            "season": season.name,
            "start": season.start_date,
            "end": season.end_date,
            "cutoff": cutoff,
            "charges": charges,
            "produits": produits,
            "totalCharges": total_charges,
            "totalProduits": total_produits,
            "result": round(total_produits - total_charges, 2),
            "treasury": treasury,
            "treasuryTotal": round(sum(t["amount"] for t in treasury), 2),
            "unpaid": unpaid,
            "unpaidTotal": round(sum(u["amount"] for u in unpaid), 2),
            "events": sorted(events.values(), key=lambda e: e["event"]),
            "unclassified": round(sum(a.balance for a in accounts if a.number == UNCLASSIFIED[EXPENSE][0]) / 100
                                  - sum(a.balance for a in accounts if a.number == UNCLASSIFIED[INCOME][0]) / 100, 2),
        }

    def synthese_xlsx(self, synthese: dict, org_name: str) -> bytes:
        bold = Font(bold=True)
        head = PatternFill("solid", fgColor="E9E3F5")
        money = '#,##0.00 "€";-#,##0.00 "€";""'
        wb = Workbook()
        ws = wb.active
        ws.title = "Synthèse"
        ws.append([f"{org_name} — Synthèse Dépenses-Recettes {synthese['season']}"])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([f"Du {_fr(synthese['start'])} au {_fr(synthese['cutoff'])}"])
        ws.append([])

        def block(title: str, groups: list[dict], total_label: str, total: float) -> None:
            ws.append([title])
            ws.cell(ws.max_row, 1).font, ws.cell(ws.max_row, 1).fill = bold, head
            for g in groups:
                ws.append([f"{g['code']} · {g['label']}"])
                ws.cell(ws.max_row, 1).font = bold
                for line in g["lines"]:
                    ws.append([f"{line['number']} {line['label']}", line["amount"]])
                    ws.cell(ws.max_row, 2).number_format = money
            ws.append([total_label, total])
            ws.cell(ws.max_row, 1).font = ws.cell(ws.max_row, 2).font = bold
            ws.cell(ws.max_row, 2).number_format = money
            ws.append([])

        block("CHARGES", synthese["charges"], "Total des charges", synthese["totalCharges"])
        block("PRODUITS", synthese["produits"], "Total des produits", synthese["totalProduits"])
        ws.append(["Déficit" if synthese["result"] < 0 else "Excédent", synthese["result"]])
        ws.cell(ws.max_row, 1).font = ws.cell(ws.max_row, 2).font = bold
        ws.cell(ws.max_row, 2).number_format = money
        ws.append([])
        ws.append([f"Trésorerie au {_fr(synthese['cutoff'])}"])
        ws.cell(ws.max_row, 1).font, ws.cell(ws.max_row, 1).fill = bold, head
        for t in synthese["treasury"]:
            ws.append([t["label"], t["amount"]])
            ws.cell(ws.max_row, 2).number_format = money
        ws.append(["Solde de trésorerie", synthese["treasuryTotal"]])
        ws.cell(ws.max_row, 1).font = ws.cell(ws.max_row, 2).font = bold
        ws.cell(ws.max_row, 2).number_format = money
        ws.append(["Fournisseurs à payer", synthese["unpaidTotal"]])
        ws.cell(ws.max_row, 2).number_format = money
        for e in synthese["events"]:
            ws.append([])
            ws.append([e["event"].capitalize()])
            ws.cell(ws.max_row, 1).font, ws.cell(ws.max_row, 1).fill = bold, head
            for label, key in (("Ressources", "resources"), ("Coût", "cost"), ("Bénéfice", "profit")):
                ws.append([f"{label} : {e['event']}", e[key]])
                ws.cell(ws.max_row, 2).number_format = money
        ws.column_dimensions["A"].width = 44
        ws.column_dimensions["B"].width = 16
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    # ----- Balance et Grand livre -----

    def balance_pdf(self, db: Session, member: Member, season: Season, org_name: str) -> bytes:
        books = self.books(db, member, season)
        cutoff = books.pop("_cutoff")
        e = html.escape
        rows, total_d, total_c = [], 0, 0
        for a in sorted(books.values(), key=lambda a: a.number):
            d, c = a.debit, a.credit
            total_d, total_c = total_d + d, total_c + c
            sd, sc = max(a.balance, 0), max(-a.balance, 0)
            rows.append(
                f"<tr><td>{e(a.number)}</td><td class='l'>{e(a.label)}</td><td>{_c(d)}</td><td>{_c(c)}</td><td>{_c(sd)}</td><td>{_c(sc)}</td></tr>"
            )
        result = sum(a.balance for a in books.values() if a.nature in (EXPENSE, INCOME))
        body = (
            f"<h1>{e(org_name)} — Balance {e(season.name)}</h1>"
            f"<p class='muted'>Du {_fr(season.start_date)} au {_fr(cutoff)}. "
            + ("Partie double : achats et honoraires par les comptes de tiers, lettrés avec leur paiement." if self.accounting.level(db, member.org_id) == FULL else "Comptabilité de trésorerie : chaque ligne de banque et de caisse contre son compte du plan.")
            + "</p>"
            "<table><tr><th>Compte</th><th class='l'>Libellé</th><th>Débit</th><th>Crédit</th><th>Solde débiteur</th><th>Solde créditeur</th></tr>"
            + "".join(rows)
            + f"<tr class='total'><td></td><td class='l'>Totaux</td><td>{_c(total_d)}</td><td>{_c(total_c)}</td><td></td><td></td></tr></table>"
            f"<h2>{'Déficit' if result > 0 else 'Excédent'} de la saison : {_c(abs(result))}</h2>"
        )
        return html_to_pdf(body, _CSS + " th.l, td.l { text-align: left; }")

    def grand_livre_pdf(self, db: Session, member: Member, season: Season, org_name: str) -> bytes:
        books = self.books(db, member, season)
        cutoff = books.pop("_cutoff")
        e = html.escape
        parts = [
            f"<h1>{e(org_name)} — Grand livre {e(season.name)}</h1>",
            f"<p class='muted'>Du {_fr(season.start_date)} au {_fr(cutoff)}. Journaux : AN à-nouveaux, BQn banque, CA caisse, AC achats, HO honoraires, NF notes de frais.</p>",
        ]
        total_d = total_c = 0
        for a in sorted(books.values(), key=lambda a: a.number):
            running = 0
            lines = []
            for l in a.lines:
                running += l.debit - l.credit
                lines.append(
                    f"<tr><td>{_fr(l.date)}</td><td>{e(l.journal)}</td><td class='l'>{e(l.piece)}</td><td class='l'>{e(l.label)}</td><td>{e(l.letter)}</td>"
                    f"<td>{_c(l.debit)}</td><td>{_c(l.credit)}</td><td>{_signed(running)}</td></tr>"
                )
            total_d, total_c = total_d + a.debit, total_c + a.credit
            parts.append(
                f"<h2>{e(a.number)} {e(a.label)}</h2><table class='small'>"
                "<tr><th>Date</th><th>Journal</th><th class='l'>Pièce</th><th class='l'>Libellé</th><th>Lettre</th><th>Débit</th><th>Crédit</th><th>Solde</th></tr>"
                + "".join(lines)
                + f"<tr class='total'><td></td><td></td><td></td><td class='l'>Total du compte</td><td></td><td>{_c(a.debit)}</td><td>{_c(a.credit)}</td><td>{_signed(a.balance)}</td></tr></table>"
            )
        parts.append(f"<h2>Total général : débit {_c(total_d)}, crédit {_c(total_c)}</h2>")
        return html_to_pdf("".join(parts), _CSS + " th.l, td.l { text-align: left; }")

    # ----- Honoraires par professeur -----

    def fees_xlsx(self, db: Session, member: Member, season: Season, org_name: str) -> bytes:
        fees = self.ports.fees(db, member.org_id, season.start_date, season.end_date) if self.ports.fees else []
        months = sorted({f["month"] for f in fees})
        by_person: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
        for f in fees:
            by_person[f["person"]][f["month"]] += f["amount"]
        bold = Font(bold=True)
        money = '#,##0.00 "€";-#,##0.00 "€";""'
        wb = Workbook()
        ws = wb.active
        ws.title = "Honoraires"
        ws.append([f"{org_name} — Honoraires par professeur {season.name} (DAS-2)"])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append(["Professeur", *months, "Total"])
        for cell in ws[2]:
            cell.font = bold
        for person in sorted(by_person):
            amounts = [by_person[person].get(m, 0.0) for m in months]
            ws.append([person, *amounts, round(sum(amounts), 2)])
            for col in range(2, len(months) + 3):
                ws.cell(ws.max_row, col).number_format = money
        ws.append(["Total", *[round(sum(by_person[p].get(m, 0.0) for p in by_person), 2) for m in months], round(sum(f["amount"] for f in fees), 2)])
        for cell in ws[ws.max_row]:
            cell.font = bold
        for col in range(2, len(months) + 3):
            ws.cell(ws.max_row, col).number_format = money
        ws.column_dimensions["A"].width = 30
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    # ----- FEC et son contrôle -----

    def fec_entries(self, db: Session, member: Member, season: Season) -> list[dict]:
        """Les écritures de la saison (celles du Grand livre), regroupées :
        une écriture = les lignes d'un même journal, d'un même jour, d'une
        même opération de banque (ou du même libellé ; les à-nouveaux en une
        seule), numérotées par date."""
        books = self.books(db, member, season)
        books.pop("_cutoff")
        groups: dict[tuple, list[tuple[Account, Line]]] = defaultdict(list)
        for a in books.values():
            for l in a.lines:
                same = "" if l.journal == "AN" else l.operation_id if l.operation_id is not None else l.label
                groups[(l.date, l.journal, same)].append((a, l))
        lettered: dict[tuple[str, str], str] = {}
        for a in books.values():
            for l in a.lines:
                if l.letter:
                    key = (a.number, l.letter)
                    lettered[key] = max(lettered.get(key, ""), l.date)
        entries = []
        for n, key in enumerate(sorted(groups, key=lambda k: (k[0], k[1] != "AN", k[1], str(k[2]))), start=1):
            day, journal, _ = key
            entries.append({"number": n, "date": day, "journal": journal, "lines": [
                {"account": a.number, "accountLabel": a.label, "label": l.label, "piece": l.piece, "debit": l.debit,
                 "credit": l.credit, "letter": l.letter, "letterDate": lettered.get((a.number, l.letter), "") if l.letter else ""}
                for a, l in groups[key]
            ]})
        return entries

    def fec(self, db: Session, member: Member, season: Season) -> bytes:
        def d(day: str) -> str:
            return day.replace("-", "")

        def amount(cents: int) -> str:
            return f"{cents // 100},{cents % 100:02d}"

        def clean(text: str) -> str:
            return " ".join(str(text or "").replace("\t", " ").split())

        rows = ["\t".join(FEC_COLUMNS)]
        for e in self.fec_entries(db, member, season):
            journal = e["journal"]
            for l in e["lines"]:
                rows.append("\t".join([
                    journal, JOURNALS.get(journal[:2], journal), str(e["number"]), d(e["date"]), l["account"], clean(l["accountLabel"]),
                    "", "", clean(l["piece"]) or str(e["number"]), d(e["date"]), clean(l["label"]) or JOURNALS.get(journal[:2], journal),
                    amount(l["debit"]), amount(l["credit"]), l["letter"], d(l["letterDate"]), d(e["date"]), "", "",
                ]))
        return ("\r\n".join(rows) + "\r\n").encode("iso8859_15", errors="replace")

    def fec_check(self, db: Session, member: Member, season: Season) -> dict:
        """Ce qui empêche le FEC d'être juste (spec §8.6) : rien à corriger
        quand problems est vide ; notes = ce qu'il faut savoir sans être faux."""
        entries = self.fec_entries(db, member, season)
        lines = [(e, l) for e in entries for l in e["lines"]]
        problems, notes = [], []
        unbalanced = [e["number"] for e in entries if sum(l["debit"] - l["credit"] for l in e["lines"])]
        if unbalanced:
            problems.append(f"{len(unbalanced)} écriture(s) déséquilibrée(s) : n° {', '.join(map(str, unbalanced[:10]))}.")
        total = sum(l["debit"] - l["credit"] for _, l in lines)
        if total:
            problems.append(f"Débit et crédit diffèrent de {_eur(abs(total) / 100)}.")
        unclassified = [l for _, l in lines if l["account"] in {n for n, _ in UNCLASSIFIED.values()}]
        if unclassified:
            amount = sum(l["debit"] + l["credit"] for l in unclassified)
            problems.append(f"{len(unclassified)} ligne(s) à classer ({_eur(amount / 100)}) : Comptes › Sans compte.")
        outside = [e["number"] for e in entries if not season.start_date <= e["date"] <= season.end_date]
        if outside:
            problems.append(f"{len(outside)} écriture(s) hors de la saison.")
        bad = sorted({l["account"] for _, l in lines if len(l["account"]) < 3 or l["account"][0] not in "1234567"})
        if bad:
            problems.append(f"Numéros de compte hors plan : {', '.join(bad)}.")
        nameless = sorted({l["account"] for _, l in lines if not l["accountLabel"].strip()})
        if nameless:
            problems.append(f"Comptes sans libellé : {', '.join(nameless)}.")
        open_letters: dict[tuple[str, str], int] = defaultdict(int)
        for _, l in lines:
            if l["letter"]:
                open_letters[(l["account"], l["letter"])] += l["debit"] - l["credit"]
        wrong = sorted(f"{a} {letter}" for (a, letter), rest in open_letters.items() if rest)
        if wrong:
            problems.append(f"Lettrage qui ne s'équilibre pas : {', '.join(wrong[:10])}.")
        if self.accounting.level(db, member.org_id) != FULL:
            notes.append("Niveau Simple : comptabilité de trésorerie, sans fournisseurs ni lettrage.")
        if not any(e["journal"] == "AN" for e in entries):
            notes.append("Pas d'à-nouveaux : aucun solde de départ connu pour la saison.")
        return {"ok": not problems, "entries": len(entries), "lines": len(lines), "problems": problems, "notes": notes}

    # ----- production et rangement -----

    def _build(self, db: Session, member: Member, season: Season, kind: str) -> tuple[bytes, str, str]:
        if kind not in KINDS:
            raise ValueError("Document inconnu")
        title, extension, _ = KINDS[kind]
        org_name = self.ports.org_name(db, member.org_id) if self.ports.org_name else ""
        if kind == "synthese":
            content = self.synthese_xlsx(self.synthese(db, member, season), org_name)
        elif kind == "balance":
            content = self.balance_pdf(db, member, season, org_name)
        elif kind == "grand-livre":
            content = self.grand_livre_pdf(db, member, season, org_name)
        elif kind == "fec":
            content = self.fec(db, member, season)
        else:
            content = self.fees_xlsx(db, member, season, org_name)
        name = f"{title} {season.name}.{extension}"
        media = {"pdf": "application/pdf", "txt": "text/plain; charset=iso-8859-15"}.get(
            extension, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        return content, name, media

    def produce(self, db: Session, member: Member, season: Season, kind: str) -> tuple[bytes, str, str]:
        """(contenu, nom de fichier, type) d'un document ; rangé dans Docs ›
        Financier › AG si l'appli le permet."""
        content, name, media = self._build(db, member, season, kind)
        store_file(self.ports.store, db, member, name, content)
        return content, name, media

    def in_docs(self, db: Session, member: Member, season: Season, kind: str) -> int | None:
        """Produit le document, le range dans Docs › Financier › AG et rend
        l'id du fichier (pour l'ouvrir depuis Docs)."""
        content, name, _ = self._build(db, member, season, kind)
        return store_file(self.ports.store, db, member, name, content)

    @staticmethod
    def kinds(level: str) -> list[dict]:
        """Les documents du niveau (Simple : Synthèse ; Complet : tous)."""
        return [
            {"id": key, "title": title, "format": extension, "level": needed}
            for key, (title, extension, needed) in KINDS.items()
            if level == "complet" or needed == "simple"
        ]


def _third_party(accounts: list, piece: dict | None):
    """Le compte de tiers d'une pièce (401 pour un achat, 467x pour des
    honoraires ou une note de frais, celui qui porte le nom de la personne
    s'il existe), ou None sans pièce ou sans compte de ce genre."""
    if not piece or piece.get("nature") not in THIRD_PARTY_PREFIX:
        return None
    candidates = [a for a in accounts if a.number.startswith(THIRD_PARTY_PREFIX[piece["nature"]])]
    if not candidates:
        return None
    words = [w.lower() for w in str(piece.get("person") or "").split() if len(w) > 2]
    named = [a for a in candidates if any(w in a.label.lower() for w in words)]
    return named[0] if named else min(candidates, key=lambda a: (len(a.number), a.number))


def _letter(n: int) -> str:
    """Lettre de lettrage : AAA, AAB, … (n à partir de 1)."""
    n -= 1
    return chr(65 + n // 676 % 26) + chr(65 + n // 26 % 26) + chr(65 + n % 26)


def _nature_of(number: str, db: Session, accounting: Accounting, org_id: int) -> str:
    if number.startswith(BANK_BASE):
        return "banque"
    if number == CASH:
        return "caisse"
    if number == CARRY_FORWARD:
        return "report"
    account = next((a for a in accounting.accounts(db, org_id) if a.number == number), None)
    return account.nature if account is not None else THIRD_PARTY


def store_file(store, db: Session, member: Member, name: str, content: bytes) -> int | None:
    """Range un fichier produit dans Docs › Financier › AG ; rend son id."""
    if store is None:
        return None
    import tempfile

    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / name
        path.write_bytes(content)
        ids = store(db, member, ["AG"], [(name, path)])
    return ids[0] if ids else None


def _c(cents: int) -> str:
    return "" if not cents else _eur(cents / 100)


def _signed(cents: int) -> str:
    return _eur(cents / 100) if cents else ""
