"""Export Excel HUMAIN d'une école ("Sauvegarder École", Admin > École,
2e des 2 fichiers générés — voir backup_technique.py pour le fichier
technique) — pensé pour être lu/imprimé/archivé par une personne, pas pour
être réimporté. 3 onglets : Élèves (mêmes en-têtes que "Télécharger les
nouvelles inscriptions", mais TOUS les élèves de l'école, pas seulement
les nouveaux arrivés par le parcours d'inscription en ligne), Profs (avec
relevé d'heures), Cours.
"""

from __future__ import annotations

import datetime as dt
import io
import re

from comptes import Comptes
from cours import CoursService
from eleves import Eleves
from eleves.eleves import calculer_age
from openpyxl import Workbook
from presence import Presence
from sqlalchemy.orm import Session

from .models import Ecole

# Même liste, même ordre de colonnes que inscriptions/excel_export.py
# (voir sa docstring : en-têtes EXACTS du fichier maître réel, pour rester
# fusionnable à la main). Dupliquée plutôt qu'importée : `ecoles` ne doit
# pas dépendre de `inscriptions` (sens de dépendance inverse, voir mémoire
# projet "Découpage du backend") — et cette liste est un texte fixe du
# dossier papier (spec/SPEC.md §6.5), pas une donnée qui change souvent.
NOMS_COURS = [
    "Éveil",
    "Class Ini",
    "Jazz Ini",
    "Class Moy",
    "Jazz Moy",
    "Street Moyen",
    "Jazz Junior",
    "Street Junior Inter",
    "Class Inter",
    "Jazz Inter",
    "Pointes inter",
    "Pointes AV",
    "Class AV",
    "Jazz AV",
    "Contempo Junior",
    "Contempo Inter avance",
    "Contempo Adulte",
]


def _slug(texte: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", texte).strip("_") or "Ecole"


def nom_fichier(ecole: Ecole) -> str:
    """<Nom École><Date> (demande utilisateur) — slug sûr pour un nom de
    fichier (disque, en-tête HTTP Content-Disposition)."""
    return f"{_slug(ecole.nom)}_{dt.date.today().isoformat()}.xlsx"


def _formater_heures(minutes: int) -> str:
    """750 -> "12h30" — jamais de décimal (plus lisible pour un relevé
    d'heures que "12.5")."""
    signe = "-" if minutes < 0 else ""
    minutes = abs(minutes)
    return f"{signe}{minutes // 60}h{minutes % 60:02d}"


class EcoleExport:
    def __init__(
        self, comptes: Comptes, eleves: Eleves, cours: CoursService, presence: Presence
    ) -> None:
        self.comptes = comptes
        self.eleves = eleves
        self.cours = cours
        self.presence = presence

    def generer(self, db: Session, ecole: Ecole) -> bytes:
        classeur = Workbook()
        classeur.remove(classeur.active)  # la feuille "Sheet" par défaut, vide

        self._feuille_eleves(classeur, db, self.eleves.list(db, ecole.id))
        self._feuille_profs(classeur, db, ecole.id)
        self._feuille_cours(classeur, db, ecole.id)

        tampon = io.BytesIO()
        classeur.save(tampon)
        return tampon.getvalue()

    # --- Onglet "Élèves" (même formalisme que "Télécharger les nouvelles
    # inscriptions", voir inscriptions/excel_export.py — mais TOUS les
    # élèves de l'école, pas seulement les nouveaux inscrits en ligne) ---

    def _feuille_eleves(self, classeur: Workbook, db: Session, comptes_eleves: list) -> None:
        feuille = classeur.create_sheet("Élèves")
        feuille.append(
            ["Nom adhérent", "Prénom adhérent", "Nom - Prénom parent", "E-Mail", "Adresse", "Téléphone", "Age"]
            + NOMS_COURS
        )
        for compte in comptes_eleves:
            profil = self.eleves.get_profil(db, compte.id)
            contacts = self.eleves.contacts_de_leleve(db, compte.id)
            # 1er contact enregistré comme "parent" — un élève général n'a
            # pas de champ "contact d'urgence" dédié comme une inscription
            # en ligne (voir inscriptions/excel_export.py:_ligne), les
            # contacts (§6.4) en tiennent lieu. Même ordre nom/prénom que
            # là-bas, pour rester visuellement cohérent entre les 2 tableaux.
            premier_contact = contacts[0] if contacts else None
            contact_parent = (
                " ".join(part for part in (premier_contact.nom, premier_contact.prenom) if part) or None
                if premier_contact
                else None
            )
            date_naissance = profil.date_naissance if profil else None
            if date_naissance:
                date_age = f"{date_naissance.strftime('%d/%m/%Y')} = {calculer_age(date_naissance)} ans"
            else:
                date_age = None
            noms_cours_choisis = [c.nom for c in self.cours.cours_de_leleve(db, compte.id)]
            ligne = [
                compte.nom,
                compte.prenom,
                contact_parent,
                compte.email,
                profil.adresse if profil else None,
                compte.telephone,
                date_age,
            ]
            ligne += ["X" if nom in noms_cours_choisis else None for nom in NOMS_COURS]
            feuille.append(ligne)

    # --- Onglet "Profs" (avec relevé d'heures, demande utilisateur) ---

    def _feuille_profs(self, classeur: Workbook, db: Session, ecole_id: int) -> None:
        feuille = classeur.create_sheet("Profs")
        feuille.append(
            ["Nom", "Prénom", "Email", "Téléphone", "Cours enseignés", "Heures normales", "Heures dépassement"]
        )
        for prof in self.comptes.list_par_role(db, ecole_id, role="professeur"):
            cours_noms = ", ".join(c.nom for c in self.cours.cours_du_professeur(db, prof.id))
            # Totaux tous cours confondus, depuis le début (pas de filtre
            # de période ici — voir presence.heures_professeur, `depuis`/
            # `jusqua` optionnels) : un export "photo à l'instant T" de
            # toute l'école, pas un relevé mensuel.
            heures = self.presence.heures_professeur(db, prof.id)
            minutes_normales = heures["minutes_total"] - heures["minutes_depassement"]
            feuille.append(
                [
                    prof.nom,
                    prof.prenom,
                    prof.email,
                    prof.telephone,
                    cours_noms or None,
                    _formater_heures(minutes_normales),
                    _formater_heures(heures["minutes_depassement"]),
                ]
            )

    # --- Onglet "Cours" ---

    def _feuille_cours(self, classeur: Workbook, db: Session, ecole_id: int) -> None:
        feuille = classeur.create_sheet("Cours")
        feuille.append(
            [
                "Nom",
                "Jour",
                "Heure début",
                "Heure fin",
                "Salle",
                "Professeur(s)",
                "Horaires supplémentaires",
                "Élèves inscrits",
            ]
        )
        for cours in self.cours.list(db, ecole_id):
            profs_noms = ", ".join(f"{p.prenom} {p.nom}" for p in self.cours.professeurs_du_cours(db, cours.id))
            horaires_sup = "; ".join(
                f"{h.jour} {h.heure_debut}-{h.heure_fin}" for h in cours.horaires_supplementaires
            )
            feuille.append(
                [
                    cours.nom,
                    cours.jour,
                    cours.heure_debut,
                    cours.heure_fin,
                    cours.salle,
                    profs_noms or None,
                    horaires_sup or None,
                    len(self.cours.eleves_du_cours(db, cours.id)),
                ]
            )
