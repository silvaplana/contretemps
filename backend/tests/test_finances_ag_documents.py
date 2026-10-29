"""Documents de l'AG (spec/modules/financier.md §8.6) : grand livre en
mémoire équilibré, Synthèse par classe avec trésorerie, événement et
fournisseurs à payer, Balance et Grand livre en PDF, honoraires par prof,
lignes sans compte « à classer », retrait au distributeur, bilan par compte."""

import pymupdf
import pytest
from openpyxl import load_workbook
from test_finances import COURANT, LIVRET, TRESORIER, make_statement

from finances import Accounting, AgDocuments, AgPorts, Finances, Season
from finances.report import compute_report

SEASON = Season(id=1, name="2025-2026", start_date="2025-09-01", end_date="2026-08-31")


def statements():
    return [
        (
            "c0.pdf",
            make_statement(COURANT, ("31/07/2025", 100000), [("10/08/2025", "CB AVANT SAISON", -5000, [])], "31/08/2025"),
        ),
        (
            "c1.pdf",
            make_statement(
                COURANT,
                ("31/08/2025", 95000),
                [
                    ("03/09/2025", "CB AMAZON", -3990, []),
                    ("05/09/2025", "VIR GALLI", -160000, []),
                    ("18/09/2025", "VIR HELLOASSO", 259800, []),
                    ("20/09/2025", "PRLV MYSTERE", -1200, []),
                    ("22/09/2025", "RETRAIT DAB", -20000, []),
                    ("25/09/2025", "VIR LIVRET", -10000, []),
                ],
                "30/09/2025",
            ),
        ),
        ("l1.pdf", make_statement(LIVRET, ("31/08/2025", 50000), [("26/09/2025", "VIR C/C", 10000, [])], "30/09/2025")),
    ]


@pytest.fixture
def world(db_session):
    accounting = Accounting()
    finances = Finances(accounting=accounting)
    finances.import_files(db_session, TRESORIER, statements())
    ports = AgPorts(
        unpaid=lambda db, org: [{"title": "Pesenti", "amount": 678.0}],
        fees=lambda db, org, start, end: [
            {"person": "Pesenti", "month": "2025-10", "amount": 1850.0},
            {"person": "Pesenti", "month": "2025-11", "amount": 1850.0},
            {"person": "Revelles", "month": "2025-10", "amount": 678.0},
        ],
        org_name=lambda db, org: "Contretemps",
    )
    documents = AgDocuments(ledger=finances.get_ledger, accounting=accounting, seasons=lambda db, org: [SEASON], ports=ports)
    accounts = {a["number"]: a for a in accounting.chart(db_session, TRESORIER)["accounts"]}
    accounting.update_account(db_session, TRESORIER, accounts["613"]["id"], {"number": "613", "label": "Location Galli", "event": "gala"})
    accounting.update_account(db_session, TRESORIER, accounts["707"]["id"], {"number": "707", "label": "Entrées", "event": "gala"})
    rows = {r["label"]: r for r in finances.get_ledger(db_session, TRESORIER)["rows"]}
    accounting.set_lines(db_session, TRESORIER, rows["CB AMAZON"]["id"], [{"accountId": accounts["606"]["id"]}])
    accounting.set_lines(db_session, TRESORIER, rows["VIR GALLI"]["id"], [{"accountId": accounts["613"]["id"]}])
    accounting.set_lines(
        db_session, TRESORIER, rows["VIR HELLOASSO"]["id"],
        [{"accountId": accounts["706"]["id"], "amount": 2298}, {"accountId": accounts["756"]["id"], "amount": 300}],
    )
    accounting.add_cash_line(
        db_session, TRESORIER, {"date": "2025-09-22", "label": "Retrait", "amount": 200, "transferOperationId": rows["RETRAIT DAB"]["id"]}
    )
    accounting.add_cash_line(
        db_session, TRESORIER, {"date": "2025-09-28", "label": "Entrées du gala", "amount": 150, "accountId": accounts["707"]["id"]}
    )
    return finances, accounting, documents


def test_books_are_balanced_and_carry_the_opening_treasury(world, db_session):
    _, _, documents = world
    books = documents.books(db_session, TRESORIER, SEASON)
    books.pop("_cutoff")
    assert sum(a.debit for a in books.values()) == sum(a.credit for a in books.values())
    # Report à nouveau = trésorerie au 31/08 : 950,00 (courant) + 500,00 (livret).
    assert books["110"].balance == -145000
    assert books["5121"].lines[0].debit == 95000 and books["5122"].lines[0].debit == 50000
    # Le retrait va dans la caisse, pas dans « à classer » ; le mystère y va.
    assert books["530"].balance == 20000 + 15000
    assert books["65900"].balance == 1200
    # Virement interne : d'un 512 à l'autre, aucun compte de charge.
    assert books["5122"].balance == 60000 and "VIR LIVRET" in [l.label for l in books["5122"].lines]


def test_synthese_by_class_with_event_and_treasury(world, db_session):
    _, _, documents = world
    s = documents.synthese(db_session, TRESORIER, SEASON)
    charges = {line["number"]: line["amount"] for c in s["charges"] for line in c["lines"]}
    produits = {line["number"]: line["amount"] for p in s["produits"] for line in p["lines"]}
    assert charges == {"606": 39.9, "613": 1600.0, "65900": 12.0}
    assert produits == {"706": 2298.0, "756": 300.0, "707": 150.0}
    assert (s["totalCharges"], s["totalProduits"], s["result"]) == (1651.9, 2748.0, 1096.1)
    assert [c["code"] for c in s["charges"]] == ["60", "61", "65"]
    assert s["treasuryTotal"] == pytest.approx(1450 + 1096.1 + 678 * 0)  # trésorerie = départ + résultat
    assert s["unpaid"] == [{"title": "Pesenti", "amount": 678.0}] and s["unpaidTotal"] == 678.0
    assert s["events"] == [{"event": "gala", "resources": 150.0, "cost": 1600.0, "profit": -1450.0}]
    assert s["unclassified"] == 12.0


def test_documents_are_produced(world, db_session):
    _, accounting, documents = world
    content, name, media = documents.produce(db_session, TRESORIER, SEASON, "synthese")
    assert name == "Synthèse Dépenses-Recettes 2025-2026.xlsx" and "spreadsheet" in media
    sheet = load_workbook(filename=__import__("io").BytesIO(content)).active
    assert sheet["A1"].value == "Contretemps — Synthèse Dépenses-Recettes 2025-2026"
    for kind, text in (("balance", "Balance 2025-2026"), ("grand-livre", "Grand livre 2025-2026")):
        content, name, media = documents.produce(db_session, TRESORIER, SEASON, kind)
        assert media == "application/pdf" and name.endswith(".pdf")
        page_text = "".join(p.get_text() for p in pymupdf.open(stream=content, filetype="pdf"))
        assert text in page_text and "Report à nouveau" in page_text
    content, name, _ = documents.produce(db_session, TRESORIER, SEASON, "honoraires")
    sheet = load_workbook(filename=__import__("io").BytesIO(content)).active
    assert [c.value for c in sheet[2]] == ["Professeur", "2025-10", "2025-11", "Total"]
    assert [c.value for c in sheet[3]] == ["Pesenti", 1850.0, 1850.0, 3700.0]
    with pytest.raises(ValueError):
        documents.produce(db_session, TRESORIER, SEASON, "autre")
    assert [d["id"] for d in documents.kinds("simple")] == ["synthese"]
    assert len(documents.kinds("complet")) == 4


def test_report_adds_up_by_confirmed_account(world, db_session):
    finances, _, _ = world
    report = compute_report(finances.get_ledger(db_session, TRESORIER), SEASON.start_date, SEASON.end_date)
    expense = {r["category"]: r["amount"] for r in report["expense"]}
    assert expense["Location Galli"] == 1600.0 and expense["Fournitures et petit équipement"] == 39.9
    # Sans compte : catégorie par mots-clés comme avant, à classer par l'IA.
    assert "Autres" in expense
    unclassified = [op for op in report["operations"] if op["bankCategory"] == "Autres"]
    assert {op["label"] for op in unclassified} == {"PRLV MYSTERE", "RETRAIT DAB"}
