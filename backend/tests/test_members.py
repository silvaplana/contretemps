"""Champs des membres (spec/core.md §4.3) : la fiche que l'école règle, les
réponses d'une fiche, et ce qui les suit (saison, fusion, sauvegarde)."""

import datetime as dt

import pytest
import sqlalchemy as sa
from conftest import entetes_session

from comptes import Comptes
from cours import CoursService
from ecoles import Ecoles
from eleves import Eleves
from eleves.eleves import calculer_age
from members import MemberColumn, MemberField, MemberValue
from saisons.portee import toutes_saisons
from storage import LocalStorage

pytestmark = pytest.mark.rbac_reel
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture()
def ecole(db_session):
    comptes = Comptes()
    ecole = Ecoles().create(db_session, nom="Contretemps", code_postal="83330")
    voisine = Ecoles().create(db_session, nom="Voisine", code_postal="83000")
    admin = comptes.create(db_session, ecole_id=ecole.id, role="admin", nom="Dho", prenom="Julia")
    prof = comptes.create(db_session, ecole_id=ecole.id, role="professeur", nom="Blanc", prenom="Ima")
    eleve, _ = Eleves(comptes).create(db_session, ecole.id, nom="Perrin", prenom="Léon", email="leon@test.fr")
    admin_voisin = comptes.create(db_session, ecole_id=voisine.id, role="admin", nom="V", prenom="U")
    return {"ecole": ecole, "admin": admin, "prof": prof, "eleve": eleve, "admin_voisin": admin_voisin}


def _champ(client, ecole, **corps):
    reponse = client.post(f"/ecoles/{ecole['ecole'].id}/champs", json=corps, headers=entetes_session(ecole["admin"]))
    assert reponse.status_code == 201, reponse.text
    return reponse.json()


def _repondre(client, ecole, champ, valeur):
    return client.put(
        f"/membres/{ecole['eleve'].id}/champs/{champ['id']}", json={"value": valeur},
        headers=entetes_session(ecole["admin"]),
    )


def _reponses(client, ecole):
    reponse = client.get(f"/membres/{ecole['eleve'].id}/champs", headers=entetes_session(ecole["admin"]))
    assert reponse.status_code == 200, reponse.text
    # Les champs marqués (date de naissance, santé…) : test_champs_repris.
    return {r["field"]["label"]: r["value"] for r in reponse.json() if r["field"]["key"] is None}


def test_seuls_les_admins_de_l_ecole_reglent_la_fiche(client, ecole):
    url = f"/ecoles/{ecole['ecole'].id}/champs"
    assert client.post(url, json={"label": "Taille"}, headers=entetes_session(ecole["prof"])).status_code == 403
    assert client.post(url, json={"label": "Taille"}, headers=entetes_session(ecole["admin_voisin"])).status_code == 403
    champ = _champ(client, ecole, label="  Taille  de t-shirt ", format="list", choices=["S", "M", " ", "L"])
    assert (champ["label"], champ["choices"], champ["position"]) == ("Taille de t-shirt", ["S", "M", "L"], 1)
    assert client.get(url, headers=entetes_session(ecole["prof"])).status_code == 403
    assert [c["label"] for c in client.get(url, headers=entetes_session(ecole["admin"])).json()] == [
        "Date de naissance", "Adresse", "Allergies", "Traitement médical", "Informations importantes", "Taille de t-shirt",
    ]
    reponse = client.get(f"/membres/{ecole['eleve'].id}/champs", headers=entetes_session(ecole["admin_voisin"]))
    assert reponse.status_code == 403


def test_reglages_refuses(client, ecole):
    url = f"/ecoles/{ecole['ecole'].id}/champs"
    for corps in ({"label": " "}, {"label": "X", "format": "list"}, {"label": "X", "format": "group"},
                  {"label": "X", "format": "inconnu"}):
        assert client.post(url, json=corps, headers=entetes_session(ecole["admin"])).status_code == 422
    # Ce que la personne ne voit pas, elle ne le modifie pas.
    champ = _champ(client, ecole, label="Note", member_sees=False, member_edits=True, others_see=True)
    assert (champ["member_edits"], champ["others_see"]) == (False, False)


def test_une_reponse_par_format(client, ecole):
    naissance = _champ(client, ecole, label="Date de naissance", format="date", required=True)
    image = _champ(client, ecole, label="Droit à l’image", format="yes_no")
    age = _champ(client, ecole, label="Pointure", format="number")
    taille = _champ(client, ecole, label="Taille", format="list", choices=["S", "M", "L"], multiple=True)
    tel = _champ(client, ecole, label="Téléphone du médecin", format="phone")

    assert _repondre(client, ecole, naissance, "2015-03-02").status_code == 204
    assert _repondre(client, ecole, image, True).status_code == 204
    assert _repondre(client, ecole, age, "36,5").status_code == 204
    assert _repondre(client, ecole, taille, ["L", "S"]).status_code == 204
    assert _repondre(client, ecole, tel, "04 94 00 00 00").status_code == 204
    assert _reponses(client, ecole) == {
        "Date de naissance": "2015-03-02", "Droit à l’image": True, "Pointure": 36.5, "Taille": ["S", "L"],
        "Téléphone du médecin": "04 94 00 00 00",
    }
    for champ, valeur in ((naissance, "02/03/2015"), (image, "oui"), (age, "beaucoup"), (taille, ["XL"]), (tel, "abc")):
        assert _repondre(client, ecole, champ, valeur).status_code == 422
    # Vider efface la réponse.
    assert _repondre(client, ecole, tel, "").status_code == 204
    assert _reponses(client, ecole)["Téléphone du médecin"] is None


def test_un_champ_groupe_inscrit_aux_cours(client, db_session, ecole):
    service = CoursService()
    jazz = service.create(db_session, ecole["ecole"].id, "Jazz")
    salsa = service.create(db_session, ecole["ecole"].id, "Salsa")
    choix = client.get(f"/ecoles/{ecole['ecole'].id}/champs/groupes", headers=entetes_session(ecole["admin"])).json()
    cours = next(p for p in choix["parents"] if p["name"] == "Cours")
    assert [c["name"] for c in cours["children"]] == ["Jazz", "Salsa"]
    assert not any(a["name"].startswith("Accès") for a in choix["audiences"])

    champ = _champ(client, ecole, label="Cours suivis", format="group", group_id=cours["id"], multiple=True)
    assert _repondre(client, ecole, champ, [jazz.group_id, salsa.group_id]).status_code == 204
    assert [e.id for e in service.eleves_du_cours(db_session, salsa.id)] == [ecole["eleve"].id]
    assert _repondre(client, ecole, champ, [jazz.group_id]).status_code == 204
    assert service.eleves_du_cours(db_session, salsa.id) == []
    assert _reponses(client, ecole)["Cours suivis"] == [jazz.group_id]
    # Inscrit par Admin > Cours : la réponse suit.
    service.inscrire_eleve(db_session, salsa.id, ecole["eleve"].id)
    assert _reponses(client, ecole)["Cours suivis"] == [jazz.group_id, salsa.group_id]


def test_retirer_un_champ_garde_ses_reponses(client, db_session, ecole):
    champ = _champ(client, ecole, label="Allergies", format="long_text")
    _repondre(client, ecole, champ, "Arachides")
    entetes = entetes_session(ecole["admin"])

    assert client.delete(f"/champs/{champ['id']}", headers=entetes).status_code == 204
    assert _reponses(client, ecole) == {}
    assert client.post(f"/champs/{champ['id']}/remettre", headers=entetes).status_code == 200
    assert _reponses(client, ecole) == {"Allergies": "Arachides"}
    assert client.patch(f"/champs/{champ['id']}", json={"label": "Allergies connues"}, headers=entetes).json()[
        "label"
    ] == "Allergies connues"


def test_un_fichier_en_reponse(client, ecole, tmp_path, monkeypatch):
    from app import main

    monkeypatch.setattr(main.members_receiver, "storage", LocalStorage(tmp_path, 1024 * 1024))
    champ = _champ(client, ecole, label="Certificat médical", format="file")
    url = f"/membres/{ecole['eleve'].id}/champs/{champ['id']}/fichier"
    entetes = entetes_session(ecole["admin"])

    reponse = client.post(url, files={"fichier": ("certificat.pdf", b"%PDF-1.4 test", "application/pdf")}, headers=entetes)

    assert reponse.status_code == 200, reponse.text
    assert _reponses(client, ecole)["Certificat médical"]["name"] == "certificat.pdf"
    telecharge = client.get(url, headers=entetes)
    assert (telecharge.status_code, telecharge.content) == (200, b"%PDF-1.4 test")
    assert client.get(url, headers=entetes_session(ecole["prof"])).status_code == 403


def test_la_nouvelle_saison_recopie_la_fiche_et_les_reponses(client, db_session, ecole):
    champ = _champ(client, ecole, label="Date de naissance", format="date")
    _champ(client, ecole, label="Vide")
    _repondre(client, ecole, champ, "2015-03-02")

    reponse = client.post(
        f"/ecoles/{ecole['ecole'].id}/saisons",
        json={"nom": "2027-2028", "date_debut": "2027-09-01", "date_fin": "2028-08-31",
              "dupliquer_profs": True, "dupliquer_cours": True, "dupliquer_eleves": True},
        headers=entetes_session(ecole["admin"]),
    )

    assert reponse.status_code == 201, reponse.text
    nouvelle = reponse.json()["saison"]["id"]
    with toutes_saisons(db_session):
        champs = db_session.scalars(
            sa.select(MemberField).where(MemberField.season_id == nouvelle, MemberField.key.is_(None))
        ).all()
        assert [c.label for c in champs] == ["Date de naissance", "Vide"]
        valeur = db_session.scalar(sa.select(MemberValue).where(MemberValue.field_id == champs[0].id))
        fiche = Comptes().get(db_session, valeur.member_id)
        assert (valeur.value, fiche.compte_precedent_id) == ("2015-03-02", ecole["eleve"].id)


def test_fusion_suppression_et_sauvegarde(client, db_session, ecole):
    champ = _champ(client, ecole, label="Allergies", format="long_text")
    _repondre(client, ecole, champ, "Arachides")
    ecole_id = ecole["ecole"].id
    entetes = entetes_session(ecole["admin"])

    sauvegarde = client.get(f"/ecoles/{ecole_id}/export-technique", headers=entetes).content
    reponse = client.post(
        f"/ecoles/{ecole_id}/restaurer", files={"fichier": ("b.xlsx", sauvegarde, XLSX)}, headers=entetes
    )
    assert reponse.status_code == 204, reponse.text
    db_session.expire_all()
    assert _reponses(client, ecole) == {"Allergies": "Arachides"}

    # Fusion : la réponse de la fiche absorbée passe sur la fiche gardée.
    double, _ = Eleves(Comptes()).create(db_session, ecole_id, nom="Perrin", prenom="Léo")
    tel = _champ(client, ecole, label="Téléphone du médecin", format="phone")
    assert client.put(f"/membres/{double.id}/champs/{tel['id']}", json={"value": "0494"}, headers=entetes).status_code == 204
    reponse = client.post(f"/comptes/{ecole['eleve'].id}/fusion", json={"autre_id": double.id}, headers=entetes)
    assert reponse.status_code == 200, reponse.text
    db_session.expire_all()
    assert _reponses(client, ecole) == {"Allergies": "Arachides", "Téléphone du médecin": "0494"}

    # Suppression de la fiche : ses réponses partent avec elle.
    Eleves(Comptes()).delete(db_session, ecole["eleve"].id)
    assert db_session.scalars(sa.select(MemberValue)).all() == []


def test_l_affichage_dans_admin(client, db_session, ecole):
    ecole_id = ecole["ecole"].id
    entetes = entetes_session(ecole["admin"])
    taille = _champ(client, ecole, label="Taille", format="list", choices=["S", "M"])
    _repondre(client, ecole, taille, "M")
    url = f"/ecoles/{ecole_id}/colonnes"
    assert client.get(url, headers=entetes).json() == []  # l'ordre d'origine, tout affiché

    voulu = [{"key": "nom", "shown": True}, {"key": f"field:{taille['id']}", "shown": True},
             {"key": "paiement", "shown": False}]
    assert client.put(url, json=voulu, headers=entetes).json() == voulu
    assert client.get(url, headers=entetes).json() == voulu
    assert client.put(url, json=voulu, headers=entetes_session(ecole["prof"])).status_code == 403
    assert client.put(url, json=voulu + [voulu[0]], headers=entetes).status_code == 422
    assert client.put(url, json=[{"key": "DROP TABLE", "shown": True}], headers=entetes).status_code == 422

    # Les réponses de tous les élèves d'un coup, pour le tableau.
    reponses = client.get(f"/ecoles/{ecole_id}/champs/reponses", headers=entetes).json()
    assert reponses == {str(ecole["eleve"].id): {str(taille["id"]): "M"}}

    # Nouvelle saison : même affichage, sur les champs recopiés.
    reponse = client.post(
        f"/ecoles/{ecole_id}/saisons", json={"nom": "2027-2028", "date_debut": "2027-09-01", "date_fin": "2028-08-31"},
        headers=entetes,
    )
    assert reponse.status_code == 201, reponse.text
    nouvelle = reponse.json()["saison"]["id"]
    with toutes_saisons(db_session):
        nouveau = db_session.scalar(
            sa.select(MemberField).where(MemberField.season_id == nouvelle, MemberField.key.is_(None))
        )
        colonnes = db_session.scalars(
            sa.select(MemberColumn).where(MemberColumn.season_id == nouvelle).order_by(MemberColumn.position)
        ).all()
        assert [(c.key, c.shown) for c in colonnes] == [("nom", True), (f"field:{nouveau.id}", True), ("paiement", False)]


def test_champs_repris(client, db_session, ecole):
    """Étape 4c : date de naissance, adresse et santé sont des champs de la
    fiche comme les autres ; l'école les renomme ou les retire, l'appli les
    retrouve par leur marque (âge, exports, assistant)."""
    ecole_id = ecole["ecole"].id
    entetes = entetes_session(ecole["admin"])
    eleve, _ = Eleves(Comptes()).create(
        db_session, ecole_id, nom="Martin", prenom="Ana", date_naissance=dt.date(2015, 3, 2), allergies="Pollen",
    )
    champs = {c["key"]: c for c in client.get(f"/ecoles/{ecole_id}/champs", headers=entetes).json() if c["key"]}
    assert list(champs) == ["birth_date", "address", "allergies", "medical_treatment", "important_info"]
    reponses = client.get(f"/membres/{eleve.id}/champs", headers=entetes).json()
    assert {r["field"]["label"]: r["value"] for r in reponses if r["value"]} == {
        "Date de naissance": "2015-03-02", "Allergies": "Pollen",
    }

    # Saisi dans « Nos champs » : la fiche de l'élève (et l'âge) suit.
    url = f"/membres/{eleve.id}/champs/{champs['birth_date']['id']}"
    assert client.put(url, json={"value": "2016-03-02"}, headers=entetes).status_code == 204
    corps = client.get(f"/eleves/{eleve.id}", headers=entetes).json()
    assert (corps["date_naissance"], corps["allergies"]) == ("2016-03-02", "Pollen")
    assert corps["age"] == calculer_age(dt.date(2016, 3, 2))

    # Renommé : la marque reste ; le format ne change pas (l'âge lit une date).
    champ = f"/champs/{champs['allergies']['id']}"
    assert client.patch(champ, json={"label": "Allergies connues"}, headers=entetes).status_code == 200
    assert client.patch(champ, json={"format": "short_text"}, headers=entetes).status_code == 422
    assert client.put(f"/eleves/{eleve.id}", json={"allergies": "Arachide"}, headers=entetes).status_code == 200
    reponses = client.get(f"/membres/{eleve.id}/champs", headers=entetes).json()
    assert {r["field"]["label"]: r["value"] for r in reponses}["Allergies connues"] == "Arachide"

    # Retiré : il quitte la fiche, n'est pas recréé, et passe retiré à la saison suivante.
    assert client.delete(champ, headers=entetes).status_code == 204
    assert client.put(f"/eleves/{eleve.id}", json={"allergies": "Lait"}, headers=entetes).status_code == 200
    labels = [c["label"] for c in client.get(f"/ecoles/{ecole_id}/champs", headers=entetes).json()]
    assert "Allergies" not in labels and "Allergies connues" not in labels
    reponse = client.post(
        f"/ecoles/{ecole_id}/saisons",
        json={"nom": "2027-2028", "date_debut": "2027-09-01", "date_fin": "2028-08-31",
              "dupliquer_profs": True, "dupliquer_cours": True, "dupliquer_eleves": True},
        headers=entetes,
    )
    assert reponse.status_code == 201, reponse.text
    nouvelle = reponse.json()["saison"]["id"]
    with toutes_saisons(db_session):
        repris = db_session.scalars(
            sa.select(MemberField).where(MemberField.season_id == nouvelle, MemberField.key.is_not(None))
        ).all()
        assert sorted((c.key, c.removed_at is not None) for c in repris) == sorted(
            (k, k == "allergies") for k in champs
        )


def test_une_ancienne_sauvegarde_reprend_les_champs_des_eleves(client, db_session, ecole):
    """Sauvegarde d'avant l'étape 4c : la santé dans l'onglet ProfilsEleves."""
    import io

    from openpyxl import load_workbook

    ecole_id = ecole["ecole"].id
    entetes = entetes_session(ecole["admin"])
    classeur = load_workbook(io.BytesIO(client.get(f"/ecoles/{ecole_id}/export-technique", headers=entetes).content))
    feuille = classeur["ProfilsEleves"]
    en_tetes = [c.value for c in feuille[1]]
    for ligne in feuille.iter_rows(min_row=2):
        if ligne[en_tetes.index("compte_id")].value == ecole["eleve"].id:
            ligne[en_tetes.index("allergies")].value = "Pollen"
            ligne[en_tetes.index("date_naissance")].value = dt.date(2015, 3, 2)
    for nom in ("ChampsMembres", "ColonnesMembres", "ReponsesMembres"):
        del classeur[nom]
    sortie = io.BytesIO()
    classeur.save(sortie)

    reponse = client.post(
        f"/ecoles/{ecole_id}/restaurer", files={"fichier": ("b.xlsx", sortie.getvalue(), XLSX)}, headers=entetes
    )

    assert reponse.status_code == 204, reponse.text
    db_session.expire_all()
    corps = client.get(f"/eleves/{ecole['eleve'].id}", headers=entetes).json()
    assert (corps["date_naissance"], corps["allergies"]) == ("2015-03-02", "Pollen")
