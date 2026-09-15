"""Tests de l'import du fichier élèves officiel (voir spec/SPEC.md
§6.4bis) — MÊME code que "Intégrer fichier élèves officiel" (Admin >
École) et "Importer" (Admin > Élèves)."""

import io
import os

from cours import CoursService
from ecoles import Ecoles
from openpyxl import Workbook

FIXTURE = os.path.join(
    os.path.dirname(__file__), "..", "seed_data", "eleves_demo.xlsx"
)

NOMS_COURS_REELS = [
    "Éveil",
    "Class Ini",
    "Jazz Ini",
    "Class Moy",
    "Jazz Moy",
    "Street Moyen",
    "Jazz Junior",
    "Street Junior Inter",
    "Class Inter",
    "Jazz Inter",
    "Pointes inter",
    "Pointes AV",
    "Class AV",
    "Jazz AV",
    "Contempo Junior",
    "Contempo Inter avance",
    "Contempo Adulte",
]


def _creer_ecole_avec_cours(db_session):
    ecole = Ecoles().create(db_session, nom="Contretemps", code_postal="83330")
    cours_service = CoursService()
    for nom in NOMS_COURS_REELS:
        cours_service.create(db_session, ecole_id=ecole.id, nom=nom)
    return ecole


def _previsualiser(client, ecole_id, nom_fichier, contenu, content_type="application/vnd.openxmlformats"):
    return client.post(
        "/eleves/import/previsualiser",
        params={"ecole_id": ecole_id},
        files={"fichier": (nom_fichier, contenu, content_type)},
    )


def test_previsualiser_le_fichier_demo(client, db_session):
    """80 élèves fictifs, colonnes de cours abrégées (voir §6.4bis :
    "Class Av" -> Class AV), une colonne volontairement absente du
    mapping statique ("Contempo" -> un des 3 cours Contempo)."""
    ecole = _creer_ecole_avec_cours(db_session)

    with open(FIXTURE, "rb") as f:
        reponse = _previsualiser(client, ecole.id, "eleves.xlsx", f)
    assert reponse.status_code == 200
    apercu = reponse.json()

    assert len(apercu["lignes"]) == 80
    assert apercu["colonnes_non_reconnues_globales"] == ["Contempo"]
    assert all(ligne["eleve_existant_id"] is None and ligne["creer"] for ligne in apercu["lignes"])

    # Le 1er élève a un cours reconnu via abréviation, et un cours via la
    # colonne "Contempo" (non reconnue) -> pas ajouté à cours_ids.
    premiere = apercu["lignes"][0]
    assert premiere["nom"] == "Jean"
    assert "Contempo" in premiere["colonnes_non_reconnues"]
    assert len(premiere["cours_ids"]) == 1  # "Jazz Av", "Contempo" exclu

    # Le téléphone du 1er élève a été "cassé" par Excel (voir fixture).
    assert premiere["telephone_suspect"] is True
    assert premiere["telephone"] == "661602371"


def test_colonne_non_reconnue_puis_memorisee(client, db_session):
    """Voir §6.4bis : après résolution manuelle une fois, la colonne est
    reconnue automatiquement aux imports suivants."""
    ecole = _creer_ecole_avec_cours(db_session)
    cours_contemporain = next(
        c for c in CoursService().list(db_session, ecole.id) if c.nom == "Contempo Junior"
    )

    with open(FIXTURE, "rb") as f:
        premier = _previsualiser(client, ecole.id, "eleves.xlsx", f).json()
    assert premier["colonnes_non_reconnues_globales"] == ["Contempo"]

    reponse = client.post(
        "/eleves/import/mapping-colonne",
        params={"ecole_id": ecole.id},
        json={"en_tete_excel": "Contempo", "cours_id": cours_contemporain.id},
    )
    assert reponse.status_code == 201

    with open(FIXTURE, "rb") as f:
        second = _previsualiser(client, ecole.id, "eleves.xlsx", f).json()
    assert second["colonnes_non_reconnues_globales"] == []
    premiere = second["lignes"][0]
    assert len(premiere["cours_ids"]) == 2  # "Jazz Av" + "Contempo" maintenant résolu


def test_valider_cree_les_eleves_et_les_inscrit_aux_cours(client, db_session):
    ecole = _creer_ecole_avec_cours(db_session)

    with open(FIXTURE, "rb") as f:
        apercu = _previsualiser(client, ecole.id, "eleves.xlsx", f).json()

    reponse = client.post(
        "/eleves/import/valider",
        params={"ecole_id": ecole.id},
        json={"lignes": apercu["lignes"]},
    )
    assert reponse.status_code == 200
    resultat = reponse.json()
    assert resultat == {"crees": 80, "mis_a_jour": 0, "inchanges": 0}

    reponse = client.get("/eleves", params={"ecole_id": ecole.id})
    assert len(reponse.json()) == 80
    premier = next(e for e in reponse.json() if e["nom"] == "Jean" and e["prenom"] == "Camille")
    assert premier["telephone"] == "661602371"
    # Contact parent : élève majeur dans la fixture -> pas de contact.
    assert premier["contacts"] == []


def test_reimport_ne_cree_pas_de_doublon_et_ne_detecte_aucune_difference(client, db_session):
    """Réimporter EXACTEMENT le même fichier : chaque ligne retrouve son
    élève (nom+prénom, voir §6.4bis) mais ne propose AUCUNE différence
    (le fichier n'apporte rien de nouveau) — voir
    test_difference_par_champ_est_detectee_et_appliquee_selon_le_choix
    pour le cas où une valeur a vraiment changé."""
    ecole = _creer_ecole_avec_cours(db_session)

    with open(FIXTURE, "rb") as f:
        apercu = _previsualiser(client, ecole.id, "eleves.xlsx", f).json()
    client.post(
        "/eleves/import/valider", params={"ecole_id": ecole.id}, json={"lignes": apercu["lignes"]}
    )

    with open(FIXTURE, "rb") as f:
        second_apercu = _previsualiser(client, ecole.id, "eleves.xlsx", f).json()

    assert all(ligne["eleve_existant_id"] is not None for ligne in second_apercu["lignes"])
    assert all(ligne["differences"] == [] for ligne in second_apercu["lignes"])

    resultat = client.post(
        "/eleves/import/valider",
        params={"ecole_id": ecole.id},
        json={"lignes": second_apercu["lignes"]},
    ).json()
    assert resultat == {"crees": 0, "mis_a_jour": 0, "inchanges": 80}
    # Toujours 80 élèves en base, pas de doublon créé.
    assert len(client.get("/eleves", params={"ecole_id": ecole.id}).json()) == 80


def test_decocher_creer_sur_un_nouvel_eleve_ne_lecrit_pas(client, db_session):
    ecole = _creer_ecole_avec_cours(db_session)
    with open(FIXTURE, "rb") as f:
        apercu = _previsualiser(client, ecole.id, "eleves.xlsx", f).json()

    lignes = apercu["lignes"]
    lignes[0]["creer"] = False
    resultat = client.post(
        "/eleves/import/valider", params={"ecole_id": ecole.id}, json={"lignes": lignes}
    ).json()
    assert resultat == {"crees": 79, "mis_a_jour": 0, "inchanges": 0}


def test_difference_par_champ_est_detectee_et_appliquee_selon_le_choix(client, db_session):
    """Cœur de la demande utilisateur : une différence par CHAMP (pas par
    ligne), et rien n'est écrasé tant que l'admin ne l'a pas explicitement
    choisi (`champs_a_appliquer`) — contrairement à l'ancien comportement
    qui écrasait tout silencieusement."""
    ecole = _creer_ecole_avec_cours(db_session)
    from eleves import Eleves
    from comptes import Comptes

    eleves_service = Eleves(comptes=Comptes())
    compte, _profil = eleves_service.create(
        db_session, ecole_id=ecole.id, nom="Martin", prenom="Alix",
        email="ancien@test.fr", telephone="0600000000", adresse="1 rue Ancienne",
    )

    wb = Workbook()
    ws = wb.active
    ws.append(["Nom adhérent", "Prénom adhérent", "Email", "Téléphone", "Adresse"])
    ws.append(["Martin", "Alix", "nouveau@test.fr", "0600000000", "2 rue Nouvelle"])
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)

    apercu = _previsualiser(client, ecole.id, "officiel.xlsx", tampon).json()
    ligne = apercu["lignes"][0]
    assert ligne["eleve_existant_id"] == compte.id
    champs_differents = {d["champ"] for d in ligne["differences"]}
    # email ET adresse diffèrent, téléphone identique -> pas de différence dessus.
    assert champs_differents == {"email", "adresse"}
    diff_email = next(d for d in ligne["differences"] if d["champ"] == "email")
    assert diff_email["valeur_actuelle"] == "ancien@test.fr"
    assert diff_email["valeur_fichier"] == "nouveau@test.fr"

    # L'admin choisit d'appliquer SEULEMENT l'email, pas l'adresse.
    ligne["champs_a_appliquer"] = ["email"]
    resultat = client.post(
        "/eleves/import/valider", params={"ecole_id": ecole.id}, json={"lignes": [ligne]}
    ).json()
    assert resultat == {"crees": 0, "mis_a_jour": 1, "inchanges": 0}

    eleve_maj = client.get(f"/eleves/{compte.id}").json()
    assert eleve_maj["email"] == "nouveau@test.fr"
    assert eleve_maj["adresse"] == "1 rue Ancienne"  # jamais appliqué, resté tel quel


def test_difference_cours_najoute_que_les_cours_manquants(client, db_session):
    """Voir import_excel.py:_differences — jamais de désinscription
    automatique, seulement les cours du fichier que l'élève n'a pas
    déjà."""
    ecole = _creer_ecole_avec_cours(db_session)
    cours_service = CoursService()
    from comptes import Comptes
    from eleves import Eleves

    eleves_service = Eleves(comptes=Comptes())
    compte, _profil = eleves_service.create(db_session, ecole_id=ecole.id, nom="Petit", prenom="Zoe")
    cours_eveil = next(c for c in cours_service.list(db_session, ecole.id) if c.nom == "Éveil")
    cours_service.inscrire_eleve(db_session, cours_eveil.id, compte.id)

    wb = Workbook()
    ws = wb.active
    ws.append(["Nom", "Prénom", "Éveil", "Class Ini"])
    ws.append(["Petit", "Zoe", "X", "X"])
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)

    apercu = _previsualiser(client, ecole.id, "officiel.xlsx", tampon).json()
    ligne = apercu["lignes"][0]
    diff_cours = next(d for d in ligne["differences"] if d["champ"] == "cours")
    cours_class_ini = next(c for c in cours_service.list(db_session, ecole.id) if c.nom == "Class Ini")
    assert diff_cours["valeur_fichier"] == [cours_class_ini.id]  # "Éveil" déjà suivi, exclu

    ligne["champs_a_appliquer"] = ["cours"]
    client.post("/eleves/import/valider", params={"ecole_id": ecole.id}, json={"lignes": [ligne]})
    cours_de_zoe = {c.id for c in cours_service.cours_de_leleve(db_session, compte.id)}
    assert cours_de_zoe == {cours_eveil.id, cours_class_ini.id}  # ajouté, "Éveil" toujours là


def test_deux_lignes_identiques_du_fichier_sont_signalees_et_decochees(client, db_session):
    """Bug signalé : un fichier réel avec 2 lignes identiques pour la même
    personne créait 2 fiches, ET en recréait 2 de plus à chaque réimport
    (voir test suivant). Demande utilisateur explicite : jamais fusionné
    automatiquement, juste signalé — à l'admin de choisir laquelle garder
    (voir import_excel.py:previsualiser, `doublon_fichier`)."""
    ecole = _creer_ecole_avec_cours(db_session)
    wb = Workbook()
    ws = wb.active
    ws.append(["Nom adhérent", "Prénom adhérent", "Email", "Éveil"])
    ws.append(["Richard", "Sébastien", "sebastien@test.fr", "X"])
    ws.append(["Richard", "Sébastien", "sebastien@test.fr", "X"])
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)

    apercu = _previsualiser(client, ecole.id, "doublon.xlsx", tampon).json()
    assert len(apercu["lignes"]) == 2  # jamais fusionnées, chacune reste visible
    assert all(l["doublon_fichier"] is True and l["creer"] is False for l in apercu["lignes"])

    # Non validé tel quel (rien coché) : aucune création.
    resultat = client.post(
        "/eleves/import/valider", params={"ecole_id": ecole.id}, json={"lignes": apercu["lignes"]}
    ).json()
    assert resultat == {"crees": 0, "mis_a_jour": 0, "inchanges": 0}

    # L'admin choisit explicitement laquelle garder (1re seulement).
    lignes = apercu["lignes"]
    lignes[0]["creer"] = True
    resultat = client.post(
        "/eleves/import/valider", params={"ecole_id": ecole.id}, json={"lignes": lignes}
    ).json()
    assert resultat == {"crees": 1, "mis_a_jour": 0, "inchanges": 0}
    assert len(client.get("/eleves", params={"ecole_id": ecole.id}).json()) == 1


def test_reimport_apres_resolution_dun_doublon_ne_recree_rien(client, db_session):
    """Suite du test précédent : une fois le doublon résolu par l'admin (un
    seul élève créé), réimporter le MÊME fichier (toujours 2 lignes
    identiques) ne doit PAS recréer de doublon — avant le correctif,
    l'élève créé se retrouvait "indépartageable" par nom+prénom entre les
    2 lignes du fichier et chaque ligne repartait de zéro comme "nouvel
    élève" à chaque réimport (voir `ambigu`)."""
    ecole = _creer_ecole_avec_cours(db_session)

    def fichier():
        wb = Workbook()
        ws = wb.active
        ws.append(["Nom adhérent", "Prénom adhérent"])
        ws.append(["Richard", "Sébastien"])
        ws.append(["Richard", "Sébastien"])
        tampon = io.BytesIO()
        wb.save(tampon)
        tampon.seek(0)
        return tampon

    premier = _previsualiser(client, ecole.id, "doublon.xlsx", fichier()).json()
    lignes = premier["lignes"]
    lignes[0]["creer"] = True  # l'admin résout le doublon signalé
    client.post("/eleves/import/valider", params={"ecole_id": ecole.id}, json={"lignes": lignes})
    assert len(client.get("/eleves", params={"ecole_id": ecole.id}).json()) == 1

    second = _previsualiser(client, ecole.id, "doublon.xlsx", fichier()).json()
    assert all(l["eleve_existant_id"] is not None for l in second["lignes"])  # plus "nouveau"
    resultat = client.post(
        "/eleves/import/valider", params={"ecole_id": ecole.id}, json={"lignes": second["lignes"]}
    ).json()
    assert resultat["crees"] == 0
    assert len(client.get("/eleves", params={"ecole_id": ecole.id}).json()) == 1


def test_homonyme_ambigu_deja_en_base_nest_pas_recree(client, db_session):
    """2 élèves distincts portant déjà le même nom+prénom (sans date de
    naissance pour départager) : une ligne du fichier qui les vise ne doit
    jamais être créée automatiquement (voir `ambigu`), sinon elle
    grossirait le doublon à chaque réimport."""
    ecole = _creer_ecole_avec_cours(db_session)
    from comptes import Comptes
    from eleves import Eleves

    eleves_service = Eleves(comptes=Comptes())
    eleves_service.create(db_session, ecole_id=ecole.id, nom="Martin", prenom="Alix")
    eleves_service.create(db_session, ecole_id=ecole.id, nom="Martin", prenom="Alix")

    wb = Workbook()
    ws = wb.active
    ws.append(["Nom", "Prénom"])
    ws.append(["Martin", "Alix"])
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)

    apercu = _previsualiser(client, ecole.id, "ambigu.xlsx", tampon).json()
    ligne = apercu["lignes"][0]
    assert ligne["ambigu"] is True
    assert ligne["creer"] is False
    assert ligne["eleve_existant_id"] is None

    resultat = client.post(
        "/eleves/import/valider", params={"ecole_id": ecole.id}, json={"lignes": apercu["lignes"]}
    ).json()
    assert resultat == {"crees": 0, "mis_a_jour": 0, "inchanges": 0}
    assert len(client.get("/eleves", params={"ecole_id": ecole.id}).json()) == 2


def test_premiere_colonne_doit_ressembler_a_nom(client, db_session):
    """Contrainte de format annoncée à l'utilisateur (voir menu "Intégrer
    fichier élèves officiel")."""
    ecole = _creer_ecole_avec_cours(db_session)
    wb = Workbook()
    ws = wb.active
    ws.append(["Colonne Mystere", "Prénom"])
    ws.append(["Dupont", "Marie"])
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)

    reponse = _previsualiser(client, ecole.id, "mauvais.xlsx", tampon)
    assert reponse.status_code == 400
    assert "ressembler" in reponse.json()["detail"].lower()


def test_variantes_nom_adherent_et_prenom_adherent_acceptees(client, db_session):
    """"Nom adhérent"/"Prénom adhérent" (voir excel_export.py — le fichier
    que l'appli elle-même produit) doit être importable tel quel."""
    ecole = _creer_ecole_avec_cours(db_session)
    wb = Workbook()
    ws = wb.active
    ws.append(["Nom adhérent", "Prénom adhérent", "E-Mail"])
    ws.append(["Dupont", "Marie", "marie@test.fr"])
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)

    reponse = _previsualiser(client, ecole.id, "export_app.xlsx", tampon)
    assert reponse.status_code == 200
    apercu = reponse.json()
    assert apercu["lignes"][0]["nom"] == "Dupont"
    assert apercu["lignes"][0]["email"] == "marie@test.fr"


def test_extension_non_supportee_refusee(client, db_session):
    ecole = _creer_ecole_avec_cours(db_session)
    reponse = _previsualiser(client, ecole.id, "eleves.txt", io.BytesIO(b"Nom,Prenom\nA,B\n"), "text/plain")
    assert reponse.status_code == 400


def test_import_csv(client, db_session):
    ecole = _creer_ecole_avec_cours(db_session)
    contenu = "Nom;Prénom;Email;Éveil\nDupont;Marie;marie@test.fr;X\n".encode("utf-8")
    reponse = _previsualiser(client, ecole.id, "eleves.csv", io.BytesIO(contenu), "text/csv")
    assert reponse.status_code == 200
    apercu = reponse.json()
    assert len(apercu["lignes"]) == 1
    assert apercu["lignes"][0]["nom"] == "Dupont"
    assert apercu["lignes"][0]["email"] == "marie@test.fr"
    assert len(apercu["lignes"][0]["cours_ids"]) == 1


def test_date_naissance_extraite_de_la_colonne_combinee():
    """Voir §6.4bis : "la partie avant le ' = '", l'âge écrit est ignoré."""
    from eleves.import_excel import _extraire_date_naissance

    assert _extraire_date_naissance("07/01/2006 = 20 ans") == __import__("datetime").date(
        2006, 1, 7
    )
    assert _extraire_date_naissance(None) is None


def test_telephone_texte_detecte_les_nombres_suspects():
    from eleves.import_excel import _telephone_texte

    assert _telephone_texte(661602371) == ("661602371", True)
    assert _telephone_texte("06 61 60 23 71") == ("06 61 60 23 71", False)
    assert _telephone_texte(None) == (None, False)


def test_mapping_colonne_introuvable_avant_memorisation(client, db_session):
    """Sans mapping mémorisé ni abréviation connue, une colonne inconnue
    reste non reconnue — vérifie juste l'absence de faux positif."""
    ecole = _creer_ecole_avec_cours(db_session)
    wb = Workbook()
    ws = wb.active
    ws.append(["Nom", "Prénom", "Colonne Mystere"])
    ws.append(["Dupont", "Marie", "X"])
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)

    reponse = _previsualiser(client, ecole.id, "mini.xlsx", tampon)
    assert reponse.status_code == 200
    apercu = reponse.json()
    assert apercu["colonnes_non_reconnues_globales"] == ["Colonne Mystere"]
    assert apercu["lignes"][0]["cours_ids"] == []
